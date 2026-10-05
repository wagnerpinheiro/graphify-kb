#!/usr/bin/env -S uv run --script
# /// script
# requires-python = ">=3.10"
# dependencies = [
#   "mammoth>=1.8",
#   "markdownify>=0.13",
#   "pdfplumber>=0.11",
#   "pypdfium2>=4.30",
#   "pillow>=10",
#   "openpyxl>=3.1",
#   "pyoxigraph>=0.5.11",
#   "pyshacl>=0.40",
#   "rdflib>=7",
#   "pyyaml>=6",
# ]
# ///
"""kb.py — deterministic engine of the kb-setup meta skill: graph-RAG knowledge base over a local
RDF/OWL ontology (Turtle + SHACL + SPARQL, pyoxigraph in memory). Central engine shared by every workspace.

Usage (from anywhere inside a workspace):  uv run ~/.claude/skills/kb-setup/scripts/kb.py <command> [args]
The workspace root is the nearest ancestor of the current directory containing kb/config.yaml
(override with --root or KB_ROOT). `scaffold` creates kb/config.yaml in the current directory.
`vendor` copies the engine into the workspace (scripts/kb-engine/) so people without kb-setup can consult it:
uv run scripts/kb-engine/kb.py <command>. The central engine takes precedence; the copy is the fallback.

Principles:
- paths always relative to the workspace root (portable between git and OneDrive);
- change detection by content hash, never by mtime;
- nothing binary persisted: the graph is rebuilt in memory on every command;
- docs/ is never indexed;
- core ontology (prefix kb:, English) ships with the engine; each workspace adds a domain ontology.
"""
from __future__ import annotations

import argparse
import base64
import datetime as dt
import fcntl
import getpass
import hashlib
import io
import json
import os
import re
import shutil
import socket
import subprocess
import sys
import tempfile
import unicodedata
import warnings
import zipfile
from collections import Counter, defaultdict
from pathlib import Path

import yaml

warnings.filterwarnings("ignore")

ENGINE_VERSION = "2.1.0"          # semver of the engine; workspaces record it in kb/config.yaml
EXTRACTOR_VERSION = "10"           # bump when extract-det output changes
CONVERTER_VERSION = "5"           # bump when convert output changes
ENGINE_DIR = Path(__file__).resolve().parent
SKILL_DIR = ENGINE_DIR.parent
IS_COPY = (ENGINE_DIR / "ENGINE.json").exists()   # running from a workspace copy (scripts/kb-engine/)
CENTRAL_CMD = "uv run ~/.claude/skills/kb-setup/scripts/kb.py"
INSTALL_CMD = "npx skills add wagnerpinheiro/graphify-kb --skill kb-setup -g -a claude-code"


def _bundle(rel_path: str) -> Path:
    """File shipped with the engine: beside kb.py in a workspace copy, else in the skill folder."""
    p = ENGINE_DIR / rel_path
    return p if p.exists() else SKILL_DIR / rel_path


CORE_ONTOLOGY = _bundle("assets/core/kb-core.ttl")
CORE_SHAPES = _bundle("assets/core/kb-core-shapes.ttl")
PROMPTS = _bundle("assets/templates/prompts.md")
TEMPLATES = SKILL_DIR / "assets" / "templates"   # skill only: a workspace copy cannot scaffold
_VERSION_FILE = _bundle("VERSION")
KB_SETUP_VERSION = _VERSION_FILE.read_text(encoding="utf-8").strip() if _VERSION_FILE.exists() else "unknown"


def _find_root() -> Path:
    for i, a in enumerate(sys.argv):
        if a == "--root" and i + 1 < len(sys.argv):
            return Path(sys.argv[i + 1]).expanduser().resolve()
        if a.startswith("--root="):
            return Path(a.split("=", 1)[1]).expanduser().resolve()
    if os.environ.get("KB_ROOT"):
        return Path(os.environ["KB_ROOT"]).expanduser().resolve()
    cwd = Path.cwd().resolve()
    for d in [cwd, *cwd.parents]:
        if (d / "kb" / "config.yaml").exists():
            return d
    return cwd


ROOT = _find_root()
try:  # how to call this engine in messages: a workspace copy by its relative path, else the central one
    KB_CMD = f"uv run {(ENGINE_DIR / 'kb.py').relative_to(ROOT).as_posix()}"
except ValueError:
    KB_CMD = CENTRAL_CMD
KB = ROOT / "kb"
MANIFEST_DIR = KB / "manifest"
TRIPLES_DIR = KB / "triples"
ONTOLOGY_DIR = KB / "ontology"
MAPPINGS_DIR = KB / "mappings"
QUERIES_DIR = KB / "queries"
WIKI_OUT = KB / "wiki"
LOCK_FILE = KB / ".lock"


def _read_config_raw() -> dict:
    p = KB / "config.yaml"
    if p.exists():
        try:
            return yaml.safe_load(p.read_text(encoding="utf-8")) or {}
        except yaml.YAMLError:
            return {}
    return {}


_RAW_CFG = _read_config_raw()
_ONT = _RAW_CFG.get("ontology") or {}
_WS_SLUG = re.sub(r"[^a-z0-9]+", "-", unicodedata.normalize("NFKD", ROOT.name).encode("ascii", "ignore").decode().lower()).strip("-") or "ws"
NS_KB = "http://kb.local/core#"
DOMAIN_PREFIX = _ONT.get("prefix") or "dom"
NS_DOMAIN = _ONT.get("namespace") or f"http://kb.local/{_WS_SLUG}/ont#"
NS_ID = _ONT.get("instances") or f"http://kb.local/{_WS_SLUG}/id/"
PREFIXES = {
    "kb": NS_KB,
    DOMAIN_PREFIX: NS_DOMAIN,
    "id": NS_ID,
    "rdf": "http://www.w3.org/1999/02/22-rdf-syntax-ns#",
    "rdfs": "http://www.w3.org/2000/01/rdf-schema#",
    "owl": "http://www.w3.org/2002/07/owl#",
    "xsd": "http://www.w3.org/2001/XMLSchema#",
    "dcterms": "http://purl.org/dc/terms/",
    "prov": "http://www.w3.org/ns/prov#",
    "skos": "http://www.w3.org/2004/02/skos/core#",
    "sh": "http://www.w3.org/ns/shacl#",
}
SPARQL_PREFIXES = "".join(f"PREFIX {k}: <{v}>\n" for k, v in PREFIXES.items())
TYPE_NAMESPACES = (NS_KB, NS_DOMAIN)

BINARY_EXTS = {".pdf", ".docx", ".xlsx", ".xlsm"}
DEFAULT_CONFIG = {
    "title": None,
    "language": "en",
    "review_every": "14d",
    "lock_timeout": "2h",
    "sources": {"raw": ["raw"], "wiki": ["wiki"], "never_index": ["docs", "kb", "scripts", ".claude", ".git", "graphify-out"]},
    "images": {
        "min_width": 120,
        "min_height": 80,
        "min_area": 30000,
        "repeated_min_pages": 3,
        "repeated_page_ratio": 0.3,
        "vector_threshold": 60,
        "render_scale": 2.0,
    },
    "pdf": {"strip_repeated_lines": True, "repeated_line_ratio": 0.5, "heading_size_ratio": 1.15, "frame_table_ratio": 0.5},
    "xlsx": {"max_empty_rows": 50},
    "families": {},
    "section_class": {},
    "topics": {},
    "facts": {},
    "attachments": [],
    "ask": {"text_props": {}},
}

# OneDrive/Teams conflict copies: "<name>-<MACHINE>.ext"
CONFLICT_RE = re.compile(
    r"-(DESKTOP|LAPTOP|MacBook|iMac|Mac-mini|Mac-Studio|MBP|PC)[\w-]*$|\(conflicted copy|\(cópia em conflito",
    re.I,
)


# ----------------------------------------------------------------------------- util


def die(msg: str, code: int = 1):
    print(f"ERROR: {msg}", file=sys.stderr)
    sys.exit(code)


def rel(p: Path) -> str:
    return unicodedata.normalize("NFC", p.resolve().relative_to(ROOT).as_posix())


def sha256_bytes(b: bytes) -> str:
    return hashlib.sha256(b).hexdigest()


def sha256_file(p: Path) -> str:
    h = hashlib.sha256()
    with open(p, "rb") as f:
        for chunk in iter(lambda: f.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


def sha256_text(s: str) -> str:
    return sha256_bytes(s.encode("utf-8"))


def slug(s: str, keep_dots: bool = False) -> str:
    s = unicodedata.normalize("NFKD", str(s)).encode("ascii", "ignore").decode().lower()
    s = re.sub(r"[^a-z0-9.]+" if keep_dots else r"[^a-z0-9]+", "-", s)
    return s.strip("-.") or "x"


def now_iso() -> str:
    return dt.datetime.now().astimezone().replace(microsecond=0).isoformat()


def today() -> dt.date:
    return dt.date.today()


def parse_duration(s: str | int | None, default_days: int = 14) -> dt.timedelta:
    if s is None:
        return dt.timedelta(days=default_days)
    if isinstance(s, (int, float)):
        return dt.timedelta(days=s)
    m = re.fullmatch(r"\s*(\d+)\s*([hdwm]?)\s*", str(s).lower())
    if not m:
        return dt.timedelta(days=default_days)
    n, u = int(m.group(1)), m.group(2) or "d"
    return {"h": dt.timedelta(hours=n), "d": dt.timedelta(days=n), "w": dt.timedelta(weeks=n), "m": dt.timedelta(days=30 * n)}[u]


CORE_QUERIES = _bundle("assets/queries")


def find_query(name: str) -> Path | None:
    """Saved query by name: the workspace kb/queries/ first, then the core queries shipped with the engine."""
    n = name if name.endswith(".rq") else name + ".rq"
    for d in (QUERIES_DIR, CORE_QUERIES):
        if (d / n).exists():
            return d / n
    return None


def ontology_files() -> list[Path]:
    """Core ontology (engine) + domain ontologies of the workspace (kb/ontology/*.ttl except shapes)."""
    ws = sorted(p for p in ONTOLOGY_DIR.glob("*.ttl") if "shapes" not in p.name) if ONTOLOGY_DIR.exists() else []
    return [p for p in [CORE_ONTOLOGY, *ws] if p.exists()]


def shapes_files() -> list[Path]:
    ws = sorted(ONTOLOGY_DIR.glob("*shapes*.ttl")) if ONTOLOGY_DIR.exists() else []
    return [p for p in [CORE_SHAPES, *ws] if p.exists()]


def _migrate_cfg_keys(cfg: dict) -> dict:
    """Accept v1 (Portuguese) config keys transparently."""
    sys.dont_write_bytecode = True  # keep the skill folder clean (no __pycache__)
    sys.path.insert(0, str(ENGINE_DIR))
    import core_renames as R

    for old, new in R.CONFIG_KEYS.items():
        if old in cfg and new not in cfg:
            cfg[new] = cfg.pop(old)
    for t in (cfg.get("topics") or {}).values():
        for o, n in R.TOPIC_KEYS.items():
            if o in t:
                t[n] = t.pop(o)
    for f in (cfg.get("facts") or {}).values():
        for o, n in R.FACT_KEYS.items():
            if o in f:
                f[n] = f.pop(o)
        if f.get("normalize") in R.FACT_NORMALIZE:
            f["normalize"] = R.FACT_NORMALIZE[f["normalize"]]
    for a in cfg.get("attachments") or []:
        for o, n in R.ATTACHMENT_KEYS.items():
            if o in a:
                a[n] = a.pop(o)
    return cfg


def load_config() -> dict:
    cfg = json.loads(json.dumps(DEFAULT_CONFIG))
    p = KB / "config.yaml"
    if p.exists():
        user = _migrate_cfg_keys(yaml.safe_load(p.read_text(encoding="utf-8")) or {})
        for k, v in user.items():
            if isinstance(v, dict) and isinstance(cfg.get(k), dict):
                cfg[k].update(v)
            else:
                cfg[k] = v
    return cfg


def write_if_changed(p: Path, content: str) -> bool:
    if p.exists() and p.read_text(encoding="utf-8") == content:
        return False
    p.parent.mkdir(parents=True, exist_ok=True)
    p.write_text(content, encoding="utf-8")
    return True


def split_frontmatter(text: str) -> tuple[str, str]:
    """Return (frontmatter_yaml, body). Empty frontmatter if absent."""
    if text.startswith("---\n"):
        end = text.find("\n---", 4)
        if end != -1:
            nl = text.find("\n", end + 4)
            body = text[nl + 1:] if nl != -1 else ""
            return text[4:end + 1], body
    return "", text


def parse_frontmatter(text: str) -> tuple[dict, str, str]:
    fm, body = split_frontmatter(text)
    try:
        data = yaml.safe_load(fm) if fm else {}
    except yaml.YAMLError:
        data = {}
    return (data if isinstance(data, dict) else {}), fm, body


def dump_frontmatter(d: dict) -> str:
    return "---\n" + yaml.safe_dump(d, sort_keys=False, allow_unicode=True, width=1000) + "---\n"


def is_git() -> bool:
    try:
        r = subprocess.run(["git", "-C", str(ROOT), "rev-parse", "--is-inside-work-tree"], capture_output=True, text=True)
        return r.returncode == 0 and r.stdout.strip() == "true"
    except FileNotFoundError:
        return False


def current_user_name() -> str:
    if is_git():
        r = subprocess.run(["git", "-C", str(ROOT), "config", "user.name"], capture_output=True, text=True)
        if r.stdout.strip():
            return r.stdout.strip()
    try:
        r = subprocess.run(["id", "-F"], capture_output=True, text=True)
        if r.returncode == 0 and r.stdout.strip():
            return r.stdout.strip()
    except FileNotFoundError:
        pass
    return getpass.getuser()


def to_date(v) -> dt.date | None:
    if v is None or v == "":
        return None
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    s = str(v).strip()
    m = re.match(r"D?:?(\d{4})(\d{2})(\d{2})", s)  # PDF date D:YYYYMMDD
    if m and not re.match(r"\d{4}-", s):
        try:
            return dt.date(int(m[1]), int(m[2]), int(m[3]))
        except ValueError:
            return None
    m = re.match(r"(\d{4})-(\d{2})-(\d{2})", s)
    if m:
        return dt.date(int(m[1]), int(m[2]), int(m[3]))
    m = re.match(r"(\d{1,2})/(\d{1,2})/(\d{4})", s)
    if m:
        try:
            return dt.date(int(m[3]), int(m[2]), int(m[1]))
        except ValueError:
            return None
    return None


MESES = {m: i + 1 for i, m in enumerate(
    ["janeiro", "fevereiro", "março", "abril", "maio", "junho", "julho", "agosto", "setembro", "outubro", "novembro", "dezembro"])}
MESES.update({m: i + 1 for i, m in enumerate(
    ["enero", "febrero", "marzo", "abril", "mayo", "junio", "julio", "agosto", "septiembre", "octubre", "noviembre", "diciembre"])})
DATE_RE = re.compile(
    r"\b(\d{1,2})/(\d{1,2})/(\d{4})\b|\b(\d{1,2})\s+de\s+(" + "|".join(MESES) + r")\s+de\s+(\d{4})\b", re.I)


def find_dates(text: str) -> list[tuple[dt.date, str]]:
    out = []
    for m in DATE_RE.finditer(text):
        try:
            if m.group(1):
                d = dt.date(int(m[3]), int(m[2]), int(m[1]))
            else:
                d = dt.date(int(m[6]), MESES[m[5].lower()], int(m[4]))
            out.append((d, m.group(0)))
        except ValueError:
            pass
    return out


def version_key(v: str | None) -> str:
    if not v:
        return ""
    nums = re.findall(r"\d+", str(v))
    return ".".join(f"{int(n):06d}" for n in (nums + ["0", "0", "0"])[:4])


VERSION_FILE_RES = [
    re.compile(r"(?:^|[-_ .])v(\d+(?:\.\d+)*)(?=$|[-_ .])", re.I),
    re.compile(r"(\d+)\s*[ªa]\s*[-_ ]?\s*vers[aã]o", re.I),
    re.compile(r"vers[aã]o[-_ ]*(\d+(?:\.\d+)*)", re.I),
]
VERSION_TEXT_RE = re.compile(r"vers[aã]o\s*(?:n[ºo°.]*\s*)?[:\-–]?\s*v?(\d+(?:\.\d+)+|\d+)\b", re.I)


def detect_version(stem: str, cover_text: str) -> tuple[str | None, str | None]:
    """Return (version, warning)."""
    for rx in VERSION_FILE_RES:
        m = rx.search(stem)
        if m:
            return m.group(1), None
    cands = sorted({m.group(1) for m in VERSION_TEXT_RE.finditer(cover_text)}, key=version_key)
    if len(cands) == 1:
        return cands[0], None
    if len(cands) > 1:
        return None, f"ambiguous version on the cover ({', '.join(cands)}); set `version` in the frontmatter and list it in manual_fields"
    return None, "version not found (file name and cover)"


def family_stem(stem: str) -> str:
    s = stem
    for rx in VERSION_FILE_RES:
        s = rx.sub(" ", s)
    return slug(s)


# ----------------------------------------------------------------------------- source discovery


class Source:
    """An indexable source: binary in raw/, native .md in raw/, or wiki/ note."""

    def __init__(self, kind: str, path: Path, md: Path, layer: str):
        self.kind = kind  # pdf | docx | xlsx | md | note
        self.path = path  # binary (or the .md itself)
        self.md = md
        self.layer = layer
        self.rel = rel(path) if path.exists() else unicodedata.normalize("NFC", path.relative_to(ROOT).as_posix())
        self.id = slug(self.rel)
        self.family = family_stem(unicodedata.normalize("NFC", path.stem))

    @property
    def binary_present(self) -> bool:
        return self.kind not in ("md", "note") and self.path.exists()

    @property
    def manifest_path(self) -> Path:
        return MANIFEST_DIR / f"{self.id}.json"

    def det_ttl(self) -> Path:
        return TRIPLES_DIR / "det" / f"{self.id}.ttl"

    def llm_ttl(self) -> Path:
        return TRIPLES_DIR / "llm" / f"{self.id}.ttl"

    def doc_iri(self) -> str:
        return NS_ID + ("note/" if self.layer == "wiki" else "doc/") + self.id


def _skip_dir(name: str) -> bool:
    # graphify-out/: graphify output/cache (may be written inside the scanned directory, e.g. raw/)
    return name.startswith(".") or name.endswith(".assets") or name == "graphify-out"


def walk_files(base: Path):
    for dirpath, dirnames, filenames in os.walk(base):
        dirnames[:] = sorted(d for d in dirnames if not _skip_dir(d))
        for f in sorted(filenames):
            if f.startswith(".") or f.startswith("~$"):
                continue
            yield Path(dirpath) / f


def discover(cfg: dict) -> tuple[list[Source], list[str], list[str]]:
    """Return (sources, conflict_copies, unsupported)."""
    sources: list[Source] = []
    conflicts: list[str] = []
    unsupported: list[str] = []
    never = {ROOT / d for d in cfg["sources"].get("never_index", [])}
    for root_name in cfg["sources"]["raw"]:
        base = ROOT / root_name
        if not base.exists():
            continue
        files = [p for p in walk_files(base) if not any(p.is_relative_to(n) for n in never)]
        binaries_by_stem: dict[Path, list[Path]] = defaultdict(list)
        for p in files:
            if CONFLICT_RE.search(p.stem):
                conflicts.append(rel(p))
                continue
            ext = p.suffix.lower()
            if ext in BINARY_EXTS:
                binaries_by_stem[p.parent / p.stem].append(p)
        for stem, bins in binaries_by_stem.items():
            for i, b in enumerate(sorted(bins)):
                md = stem.parent / (stem.name + ".md") if i == 0 else b.with_name(b.name + ".md")
                kind = "xlsx" if b.suffix.lower() in (".xlsx", ".xlsm") else b.suffix.lower()[1:]
                sources.append(Source(kind, b, md, "raw"))
        mds_taken = {s.md for s in sources}
        for p in files:
            if CONFLICT_RE.search(p.stem) or p.suffix.lower() in BINARY_EXTS or p.suffix.lower() == ".zip":
                continue
            if p.suffix.lower() == ".md":
                if p in mds_taken:
                    continue
                meta, _, _ = parse_frontmatter(p.read_text(encoding="utf-8", errors="replace"))
                if meta.get("source_file"):  # converted .md whose binary is absent (e.g. git clone)
                    src_path = ROOT / meta["source_file"]
                    ext = src_path.suffix.lower()
                    kind = "xlsx" if ext in (".xlsx", ".xlsm") else ext[1:]
                    s = Source(kind, src_path, p, "raw")
                    sources.append(s)
                else:
                    sources.append(Source("md", p, p, "raw"))
            else:
                unsupported.append(rel(p))
    for root_name in cfg["sources"]["wiki"]:
        base = ROOT / root_name
        if not base.exists():
            continue
        for p in walk_files(base):
            if CONFLICT_RE.search(p.stem):
                conflicts.append(rel(p))
            elif p.suffix.lower() == ".md":
                sources.append(Source("note", p, p, "wiki"))
    # conflict copies inside kb/ too
    for p in walk_files(KB):
        if CONFLICT_RE.search(p.stem):
            conflicts.append(rel(p))
    assign_families(sources, cfg)
    return sources, sorted(set(conflicts)), unsupported


def assign_families(sources: list[Source], cfg: dict):
    """Group versions of the same document: `families` config or common prefix of the normalized name."""
    explicit = {}
    for fam, paths in (cfg.get("families") or {}).items():
        for p in paths:
            explicit[unicodedata.normalize("NFC", p)] = slug(fam)
    raws = [s for s in sources if s.layer == "raw"]
    stems = sorted({s.family for s in raws}, key=len)
    for s in raws:
        if s.rel in explicit:
            s.family = explicit[s.rel]
            continue
        for cand in stems:
            if len(cand) >= 12 and cand != s.family and s.family.startswith(cand + "-"):
                s.family = cand
                break


# ----------------------------------------------------------------------------- manifest (R1)


def load_manifest(src: Source) -> dict:
    if src.manifest_path.exists():
        return json.loads(src.manifest_path.read_text(encoding="utf-8"))
    return {"source": src.rel, "id": src.id, "kind": src.kind, "layer": src.layer, "steps": {}}


def save_manifest(src: Source, m: dict):
    m.update({"source": src.rel, "id": src.id, "kind": src.kind, "layer": src.layer, "md": rel(src.md) if src.md.exists() else None})
    write_if_changed(src.manifest_path, json.dumps(m, indent=2, ensure_ascii=False, sort_keys=True) + "\n")


def md_hashes(md: Path) -> tuple[str, str]:
    text = md.read_text(encoding="utf-8")
    fm, body = split_frontmatter(text)
    return sha256_text(body), sha256_text(fm)


# ----------------------------------------------------------------------------- curator lock (R4)


def lock_identity() -> dict:
    return {"user": getpass.getuser(), "name": current_user_name(), "host": socket.gethostname().split(".")[0]}


def read_lock() -> dict | None:
    if LOCK_FILE.exists():
        try:
            return json.loads(LOCK_FILE.read_text(encoding="utf-8"))
        except json.JSONDecodeError:
            return {"user": "?", "host": "?", "started": "1970-01-01T00:00:00+00:00"}
    return None


def lock_is_mine(lk: dict) -> bool:
    me = lock_identity()
    return lk.get("user") == me["user"] and lk.get("host") == me["host"]


def lock_age(lk: dict) -> dt.timedelta:
    try:
        return dt.datetime.now().astimezone() - dt.datetime.fromisoformat(lk["started"])
    except Exception:
        return dt.timedelta(days=999)


class CuratorLock:
    """Advisory lock in kb/.lock. Reentrant for the same user/machine."""

    def __init__(self, cfg: dict, command: str):
        self.cfg, self.command, self.acquired = cfg, command, False

    def __enter__(self):
        lk = read_lock()
        if lk and not lock_is_mine(lk):
            age = lock_age(lk)
            msg = (f"KB locked by {lk.get('name') or lk.get('user')} ({lk.get('user')}@{lk.get('host')}) "
                   f"since {lk.get('started')} [command: {lk.get('command')}].")
            if age > parse_duration(self.cfg.get("lock_timeout"), 0):
                msg += f" The lock is {age} old (over lock_timeout); if the holder confirmed they are done: {KB_CMD} unlock --force"
            else:
                msg += " Coordinate on the Teams channel before continuing."
            die(msg, 2)
        if not lk:
            LOCK_FILE.parent.mkdir(parents=True, exist_ok=True)
            LOCK_FILE.write_text(json.dumps({**lock_identity(), "started": now_iso(), "command": self.command}, ensure_ascii=False) + "\n")
            self.acquired = True
        return self

    def __exit__(self, *exc):
        if self.acquired and LOCK_FILE.exists():
            LOCK_FILE.unlink()


# ----------------------------------------------------------------------------- images: markers


IMG_MARKER_RE = re.compile(r"^<!-- IMG: (?P<asset>.+?) \| (?P<loc>.+?) \| (?P<status>PENDENTE|DESCRITA|IGNORADA)(?P<note>[^>]*?)-->\s*$")


def parse_image_blocks(body: str) -> dict[str, dict]:
    """asset -> {status, desc: ['> ...' lines]} from an existing .md."""
    out = {}
    lines = body.splitlines()
    for i, ln in enumerate(lines):
        m = IMG_MARKER_RE.match(ln)
        if not m:
            continue
        desc = []
        j = i + 1
        while j < len(lines) and lines[j].startswith(">"):
            desc.append(lines[j])
            j += 1
        out[nfc(m["asset"])] = {"status": m["status"], "desc": desc, "loc": m["loc"], "note": m["note"].strip()}
    return out


def nfc(s: str) -> str:
    return unicodedata.normalize("NFC", s)


def marker(asset: str, loc: str, status: str = "PENDENTE", note: str = "") -> str:
    asset = nfc(asset)  # names from macOS arrive in NFD; markers and comparisons always use NFC
    return f"<!-- IMG: {asset} | {loc} | {status}{(' ' + note) if note else ''} -->"


def image_block(asset: str, loc: str, old: dict | None, reuse: dict | None) -> str:
    if old and old["status"] in ("DESCRITA", "IGNORADA"):
        return "\n".join([marker(asset, loc, old["status"], old.get("note", ""))] + old["desc"])
    if reuse:
        desc = [re.sub(r"^> \[Figura[^\]]*\]", f"> [Figura {loc_label(loc)}]", reuse["desc"][0])] + reuse["desc"][1:] if reuse["desc"] else []
        return "\n".join([marker(asset, loc, reuse["status"], reuse.get("note", ""))] + desc)
    return marker(asset, loc)


def loc_label(loc: str) -> str:
    m = re.match(r"page (\d+)", loc)
    if m:
        return f"p.{m[1]}"
    m = re.match(r"fig (\d+)", loc)
    return m[1] if m else loc


def collect_descriptions(sources: list[Source]) -> dict[str, dict]:
    """image sha -> described block (reused across identical documents)."""
    out = {}
    for s in sources:
        if not s.md.exists() or s.kind in ("md", "note"):
            continue
        m = load_manifest(s)
        imgs = m.get("images", {})
        _, body = split_frontmatter(s.md.read_text(encoding="utf-8"))
        for asset, blk in parse_image_blocks(body).items():
            if blk["status"] != "PENDENTE" and asset in imgs:
                out.setdefault(imgs[asset], blk)
    return out


def propagate_descriptions(sources: list[Source]) -> int:
    """Copy descriptions/ignored status to PENDENTE markers of identical images (same sha) in any .md."""
    known = collect_descriptions(sources)
    n = 0
    for s in sources:
        if not s.md.exists() or s.kind in ("md", "note"):
            continue
        imgs = load_manifest(s).get("images", {})
        _, body = split_frontmatter(s.md.read_text(encoding="utf-8"))
        for asset, blk in parse_image_blocks(body).items():
            src_blk = known.get(imgs.get(asset, ""))
            if blk["status"] == "PENDENTE" and src_blk:
                set_image_in(s.md, asset, src_blk["status"], src_blk["desc"], src_blk.get("note", ""))
                n += 1
    return n


def save_png(pil, path: Path) -> str:
    path.parent.mkdir(parents=True, exist_ok=True)
    if pil.mode not in ("RGB", "RGBA", "L", "LA", "P"):
        pil = pil.convert("RGB")
    buf = io.BytesIO()
    pil.save(buf, format="PNG", optimize=False)
    data = buf.getvalue()
    if not path.exists() or path.read_bytes() != data:
        path.write_bytes(data)
    return sha256_bytes(data)


# ----------------------------------------------------------------------------- conversion: PDF


def md_cell(v) -> str:
    if v is None:
        return ""
    s = str(v).replace("\r", "").strip()
    s = re.sub(r"\s*\n\s*", " ", s)
    return s.replace("|", "\\|")


def merge_wrapped_rows(rows: list[list[str]]) -> list[list[str]]:
    """PDF: wrapped cells without a horizontal border become sparse rows; merge them into the previous row."""
    n = max(len(r) for r in rows)
    if n < 3:
        return rows
    out = [rows[0]]
    for r in rows[1:]:
        filled = [i for i, c in enumerate(r) if c]
        if len(out) > 1 and filled and len(filled) <= max(1, n // 3):
            prev = out[-1]
            for i in filled:
                prev[i] = (prev[i] + " " + r[i]).strip() if i < len(prev) else r[i]
        else:
            out.append(list(r))
    return out


def md_table(rows: list[list], merge_wrapped: bool = False) -> str:
    rows = [[md_cell(c) for c in r] for r in rows if any(c not in (None, "") for c in r)]
    if not rows:
        return ""
    n = max(len(r) for r in rows)
    rows = [r + [""] * (n - len(r)) for r in rows]
    if merge_wrapped:
        rows = merge_wrapped_rows(rows)
    out = ["| " + " | ".join(rows[0]) + " |", "|" + "---|" * n]
    out += ["| " + " | ".join(r) + " |" for r in rows[1:]]
    return "\n".join(out)


def _inside(obj, boxes, key_x="x0", key_top="top") -> bool:
    cx = (obj["x0"] + obj["x1"]) / 2
    cy = (obj["top"] + obj["bottom"]) / 2
    return any(b[0] - 1 <= cx <= b[2] + 1 and b[1] - 1 <= cy <= b[3] + 1 for b in boxes)


NUMBERED_HEADING_RE = re.compile(r"^(\d+(?:\.\d+)*)\.?\s+\S")


def convert_pdf(src: Source, cfg: dict, old_imgs: dict, reuse: dict) -> tuple[str, dict, dict, list[str]]:
    """Return (md_body, meta, images{asset:sha}, warnings)."""
    import pdfplumber
    import pypdfium2 as pdfium
    import pypdfium2.raw as pdfium_c

    icfg, pcfg = cfg["images"], cfg["pdf"]
    assets_dir = assets_dir_of(src.md)
    asset_prefix = nfc(assets_dir.name)
    warns: list[str] = []

    with pdfplumber.open(str(src.path)) as pdf:
        meta_pdf = pdf.metadata or {}
        npages = len(pdf.pages)
        pages_data = []
        size_counter: Counter = Counter()
        for pg in pdf.pages:
            # "tables" covering more than half of the page are layout frames (header/footer
            # drawn as a grid) that would swallow all the text into one cell: discarded.
            page_area = float(pg.width * pg.height)
            tables = [t for t in pg.find_tables()
                      if (t.bbox[2] - t.bbox[0]) * (t.bbox[3] - t.bbox[1]) < pcfg.get("frame_table_ratio", 0.5) * page_area]
            boxes = [t.bbox for t in tables]
            filtered = pg.filter(lambda o, b=boxes: o.get("object_type") != "char" or not _inside(o, b))
            lines = filtered.extract_text_lines(return_chars=True, strip=True)
            items = []
            for ln in lines:
                sizes = [round(c["size"], 1) for c in ln["chars"] if c["text"].strip()]
                if not sizes:
                    continue
                size = Counter(sizes).most_common(1)[0][0]
                bold = all("bold" in (c.get("fontname") or "").lower() for c in ln["chars"] if c["text"].strip())
                size_counter.update(sizes)
                items.append({"t": "line", "top": ln["top"], "text": ln["text"], "size": size, "bold": bold})
            for t in tables:
                try:
                    rows = t.extract()
                except Exception:
                    rows = []
                items.append({"t": "table", "top": t.bbox[1], "rows": rows})
            vec = sum(1 for o in (pg.lines + pg.rects + pg.curves) if not _inside(o, boxes))
            pages_data.append({"items": items, "vec": vec, "height": pg.height})
        cover_text = "\n".join(i["text"] for i in pages_data[0]["items"] if i["t"] == "line") if pages_data else ""

    body_size = size_counter.most_common(1)[0][0] if size_counter else 10
    # repeated lines (header/footer)
    repeated: set[str] = set()
    if pcfg.get("strip_repeated_lines") and npages >= 3:
        cnt: Counter = Counter()
        for pd_ in pages_data:
            cnt.update({re.sub(r"\d+", "#", i["text"]) for i in pd_["items"] if i["t"] == "line"})
        thr = max(3, pcfg.get("repeated_line_ratio", 0.5) * npages)
        repeated = {k for k, v in cnt.items() if v >= thr}

    # raster images (pypdfium2)
    pdfdoc = pdfium.PdfDocument(str(src.path))
    raw_imgs: dict[int, list] = defaultdict(list)
    sha_pages: dict[str, set] = defaultdict(set)
    for i in range(len(pdfdoc)):
        page = pdfdoc[i]
        ph = page.get_height()
        for obj in page.get_objects(filter=[pdfium_c.FPDF_PAGEOBJ_IMAGE]):
            try:
                w, h = obj.get_px_size()
            except Exception:
                continue
            if w < icfg["min_width"] or h < icfg["min_height"] or w * h < icfg["min_area"]:
                continue
            try:
                pil = obj.get_bitmap(render=False).to_pil()
            except Exception:
                try:
                    pil = obj.get_bitmap(render=True).to_pil()
                except Exception:
                    continue
            key = sha256_bytes(pil.tobytes() + f"{pil.size}{pil.mode}".encode())
            try:
                top = ph - obj.get_pos()[3]
            except Exception:
                top = 0
            raw_imgs[i].append((top, pil, key))
            sha_pages[key].add(i)
    decor_thr = max(icfg["repeated_min_pages"], icfg["repeated_page_ratio"] * npages)
    decor = {k for k, pgs in sha_pages.items() if len(pgs) >= decor_thr}

    images: dict[str, str] = {}
    out: list[str] = []
    for i, pd_ in enumerate(pages_data):
        pno = i + 1
        out.append(f"\n<!-- page: {pno} -->\n")
        items = list(pd_["items"])
        rendered = pd_["vec"] >= icfg["vector_threshold"]
        if rendered:
            name = f"p{pno:03d}-render.png"
            pil = pdfdoc[i].render(scale=icfg["render_scale"]).to_pil()
            sha = save_png(pil, assets_dir / name)
            asset = f"{asset_prefix}/{name}"
            images[asset] = sha
            items.append({"t": "img", "top": -1, "asset": asset, "loc": f"page {pno}", "sha": sha})
        else:
            k = 0
            for top, pil, key in sorted(raw_imgs.get(i, []), key=lambda x: x[0]):
                if key in decor:
                    continue
                k += 1
                name = f"p{pno:03d}-{k}.png"
                sha = save_png(pil, assets_dir / name)
                asset = f"{asset_prefix}/{name}"
                images[asset] = sha
                items.append({"t": "img", "top": top, "asset": asset, "loc": f"page {pno}", "sha": sha})
        last_size = None
        for it in sorted(items, key=lambda x: x["top"]):
            if it["t"] != "line":
                last_size = None
            if it["t"] == "line":
                text = it["text"]
                if re.sub(r"\d+", "#", text) in repeated:
                    continue
                is_big = it["size"] >= body_size * pcfg["heading_size_ratio"]
                num = NUMBERED_HEADING_RE.match(text)
                if len(text) <= 150 and (is_big or (it["bold"] and num and len(text) <= 120 and not text.endswith((".", ";", ",")))):
                    if num:
                        level = min(num.group(1).count(".") + 2, 6)
                    else:
                        level = 1 if it["size"] >= body_size * 1.6 else 2
                    prev = out[-1] if out else ""
                    pm = re.match(r"\n(#+) (.*)\n$", prev)
                    if pm and len(pm.group(1)) == level and not num and last_size == it["size"]:
                        # heading wrapped over several lines (same level and size, no numbering)
                        out[-1] = f"\n{'#' * level} {pm.group(2)} {text}\n"
                    else:
                        out.append(f"\n{'#' * level} {text}\n")
                    last_size = it["size"]
                    continue
                else:
                    last_size = None
                    out.append(re.sub(r"^[•●▪■◦]\s*", "- ", text))
            elif it["t"] == "table":
                tbl = md_table(it["rows"], merge_wrapped=True)
                if tbl:
                    out.append("\n" + tbl + "\n")
            else:
                out.append("\n" + image_block(it["asset"], it["loc"], old_imgs.get(it["asset"]), reuse.get(it["sha"])) + "\n")
    pdfdoc.close()
    clean_assets(assets_dir, images)

    meta = {
        "title": (meta_pdf.get("Title") or "").strip() or src.path.stem,
        "document_date": to_date(meta_pdf.get("ModDate")) or to_date(meta_pdf.get("CreationDate")),
        "cover_text": cover_text,
    }
    if not meta["document_date"]:
        ds = find_dates(cover_text)
        meta["document_date"] = ds[0][0] if ds else None
    body = re.sub(r"\n{3,}", "\n\n", "\n".join(out)).strip() + "\n"
    return body, meta, images, warns


def assets_dir_of(md: Path) -> Path:
    return md.parent / (md.stem + ".assets")  # don't use with_suffix: names containing dots (T.I, v1.1.0)


def clean_assets(assets_dir: Path, keep: dict):
    if not assets_dir.exists():
        return
    names = {nfc(Path(a).name) for a in keep}
    for f in assets_dir.iterdir():
        if f.is_file() and nfc(f.name) not in names:
            f.unlink()
    if not any(assets_dir.iterdir()):
        assets_dir.rmdir()


# ----------------------------------------------------------------------------- conversion: DOCX


def docx_core_props(path: Path) -> dict:
    out = {}
    with zipfile.ZipFile(path) as z:
        if "docProps/core.xml" in z.namelist():
            xml = z.read("docProps/core.xml").decode("utf-8", "replace")
            for tag in ("title", "modified", "created"):
                m = re.search(rf"<(?:dc|dcterms):{tag}[^>]*>([^<]*)<", xml)
                if m:
                    out[tag] = m.group(1).strip()
    return out


def docx_images_in_order(path: Path) -> list[tuple[str, bytes]]:
    with zipfile.ZipFile(path) as z:
        doc = z.read("word/document.xml").decode("utf-8", "replace")
        rels = z.read("word/_rels/document.xml.rels").decode("utf-8", "replace") if "word/_rels/document.xml.rels" in z.namelist() else ""
        relmap = {}
        for m in re.finditer(r"<Relationship\b[^>]*>", rels):
            tag = m.group(0)
            rid = re.search(r'Id="([^"]+)"', tag)
            tgt = re.search(r'Target="([^"]+)"', tag)
            if rid and tgt and "image" in tag:
                relmap[rid.group(1)] = "word/" + tgt.group(1).lstrip("/").removeprefix("word/")
        out = []
        for m in re.finditer(r'<a:blip\b[^>]*r:embed="([^"]+)"|<v:imagedata\b[^>]*r:id="([^"]+)"', doc):
            target = relmap.get(m.group(1) or m.group(2))
            if target and target in z.namelist():
                out.append((target, z.read(target)))
        return out


def convert_docx(src: Source, cfg: dict, old_imgs: dict, reuse: dict) -> tuple[str, dict, dict, list[str]]:
    # same pipeline as markitdown's DOCX converter (mammoth -> HTML -> markdownify), without its
    # heavy dependencies (magika/onnxruntime have no wheel for Python 3.14)
    import mammoth
    from markdownify import markdownify
    from PIL import Image

    icfg = cfg["images"]
    warns: list[str] = []
    with open(src.path, "rb") as f:
        html = mammoth.convert_to_html(f, convert_image=mammoth.images.img_element(lambda image: {"src": "kbimg:"})).value
    text = markdownify(html, heading_style="ATX", bullets="-")
    assets_dir = assets_dir_of(src.md)
    imgs = docx_images_in_order(src.path)
    sha_count = Counter(sha256_bytes(b) for _, b in imgs)
    blocks: list[str | None] = []
    images: dict[str, str] = {}
    k = 0
    for target, data in imgs:
        raw_sha = sha256_bytes(data)
        try:
            pil = Image.open(io.BytesIO(data))
            pil.load()
        except Exception:
            k += 1
            name = f"img-{k:03d}{Path(target).suffix.lower()}"
            (assets_dir).mkdir(parents=True, exist_ok=True)
            (assets_dir / name).write_bytes(data)
            asset = f"{nfc(assets_dir.name)}/{name}"
            images[asset] = raw_sha
            blocks.append(marker(asset, f"fig {k}", "IGNORADA", f"(formato {Path(target).suffix} não suportado)"))
            continue
        w, h = pil.size
        if w < icfg["min_width"] or h < icfg["min_height"] or w * h < icfg["min_area"] or sha_count[raw_sha] >= icfg["repeated_min_pages"]:
            blocks.append(None)
            continue
        k += 1
        name = f"img-{k:03d}.png"
        sha = save_png(pil, assets_dir / name)
        asset = f"{nfc(assets_dir.name)}/{name}"
        images[asset] = sha
        blocks.append(image_block(asset, f"fig {k}", old_imgs.get(asset), reuse.get(sha)))
    placeholder = re.compile(r"!\[[^\]]*\]\(kbimg:\)")
    n_ph = len(placeholder.findall(text))
    if n_ph == len(blocks):
        it = iter(blocks)
        text = placeholder.sub(lambda m: ("\n\n" + b + "\n\n") if (b := next(it)) else "", text)
    else:
        if n_ph:
            warns.append(f"docx: {n_ph} images in the text × {len(blocks)} in the XML; markers moved to the Figuras section")
        text = placeholder.sub("", text)
        kept = [b for b in blocks if b]
        if kept:
            text += "\n\n## Figuras\n\n" + "\n\n".join(kept) + "\n"
    clean_assets(assets_dir, images)
    props = docx_core_props(src.path)
    cover = text[:3000]
    meta = {
        "title": props.get("title") or src.path.stem,
        "document_date": to_date(props.get("modified")) or to_date(props.get("created")),
        "cover_text": cover,
    }
    body = re.sub(r"\n{3,}", "\n\n", text).strip() + "\n"
    return body, meta, images, warns


# ----------------------------------------------------------------------------- conversion: XLSX


def xlsx_rows(ws, max_empty: int) -> list[tuple[int, list]]:
    """Non-empty rows (1-based number, values), stopping after `max_empty` consecutive empty rows."""
    out, empty = [], 0
    for idx, row in enumerate(ws.iter_rows(min_row=1, min_col=1, values_only=True), start=1):
        vals = list(row)
        if any(v not in (None, "") and str(v).strip() != "" for v in vals):
            out.append((idx, vals))
            empty = 0
        else:
            empty += 1
            if empty > max_empty:
                break
    if out:
        ncols = max(max((i for i, v in enumerate(vals) if v not in (None, "") and str(v).strip()), default=-1) + 1 for _, vals in out)
        out = [(i, (vals + [None] * ncols)[:ncols]) for i, vals in out]
    return out


def cell_str(v) -> str:
    if v is None:
        return ""
    if isinstance(v, dt.datetime):
        return v.date().isoformat() if v.time() == dt.time(0) else v.isoformat(sep=" ")
    if isinstance(v, dt.date):
        return v.isoformat()
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def detect_header(rows: list[tuple[int, list]]) -> int | None:
    if not rows:
        return None
    ncols = len(rows[0][1])
    for idx, vals in rows[:30]:
        filled = sum(1 for v in vals if cell_str(v))
        if filled >= max(2, 0.5 * ncols) and all(not isinstance(v, (int, float)) for v in vals if v is not None):
            return idx
    return rows[0][0]


def convert_xlsx(src: Source, cfg: dict, old_imgs: dict, reuse: dict) -> tuple[str, dict, dict, list[str]]:
    import openpyxl

    wb = openpyxl.load_workbook(str(src.path), read_only=True, data_only=True)
    out = []
    for ws in wb.worksheets:
        rows = xlsx_rows(ws, cfg["xlsx"]["max_empty_rows"])
        hidden = getattr(ws, "sheet_state", "visible") != "visible"
        out.append(f"\n## Aba: {ws.title}{' (oculta)' if hidden else ''}\n")
        if not rows:
            out.append("_(vazia)_")
            continue
        hdr = detect_header(rows)
        pre = [r for r in rows if r[0] < hdr]
        for idx, vals in pre:
            out.append(f"- linha {idx}: " + " | ".join(cell_str(v) for v in vals if cell_str(v)))
        table = []
        for idx, vals in rows:
            if idx == hdr:
                table.insert(0, ["Linha"] + [cell_str(v) for v in vals])
            elif idx > hdr:
                table.append([str(idx)] + [cell_str(v) for v in vals])
        if pre:
            out.append("")
        out.append(md_table(table))
    props = wb.properties
    meta = {
        "title": (props.title or "").strip() or src.path.stem,
        "document_date": to_date(props.modified) or to_date(props.created),
        "cover_text": "",
    }
    wb.close()
    body = re.sub(r"\n{3,}", "\n\n", "\n".join(out)).strip() + "\n"
    return body, meta, {}, []


CONVERTERS = {"pdf": convert_pdf, "docx": convert_docx, "xlsx": convert_xlsx}


def convert_source(src: Source, cfg: dict, reuse: dict, force: bool = False) -> str:
    """Convert a binary to .md. Return 'converted' | 'unchanged' | reason."""
    if src.kind not in CONVERTERS:
        return "n/a"
    if not src.binary_present:
        return "binary missing"
    m = load_manifest(src)
    bin_sha = sha256_file(src.path)
    conv_input = f"{bin_sha}:{CONVERTER_VERSION}"
    st = m["steps"].get("convert", {})
    if not force and st.get("input") == conv_input and src.md.exists():
        return "unchanged"
    old_meta, old_body = {}, ""
    if src.md.exists():
        old_meta, _, old_body = parse_frontmatter(src.md.read_text(encoding="utf-8"))
    old_imgs = parse_image_blocks(old_body)
    body, meta, images, warns = CONVERTERS[src.kind](src, cfg, old_imgs, reuse)
    import html as _html

    meta["title"] = unicodedata.normalize("NFC", _html.unescape(str(meta["title"]))).strip()
    version, vwarn = detect_version(unicodedata.normalize("NFC", src.path.stem), meta.pop("cover_text", ""))
    if vwarn:
        warns.append(vwarn)
    fm = {
        "title": meta["title"],
        "source_file": src.rel,
        "layer": "raw",
        "document_date": meta["document_date"].isoformat() if meta["document_date"] else None,
        "version": version,
        "converted_at": now_iso(),
    }
    if not fm["document_date"]:
        warns.append("document date not found; set `document_date` and list it in manual_fields")
    manual = old_meta.get("manual_fields") or []
    for f in manual:
        if f in old_meta:
            fm[f] = old_meta[f]
    if manual:
        fm["manual_fields"] = manual
    if warns:
        fm["warnings"] = warns
    if src.md.exists() and old_body == body:
        # same content: keep the frontmatter (avoids OneDrive/git churn) unless metadata changed
        keep = {k: v for k, v in fm.items() if k != "converted_at"}
        old_cmp = {k: v for k, v in old_meta.items() if k != "converted_at"}
        if keep == old_cmp:
            fm["converted_at"] = old_meta.get("converted_at", fm["converted_at"])
    write_if_changed(src.md, dump_frontmatter(fm) + "\n" + body)
    m["sha256"] = bin_sha
    m["images"] = images
    m["steps"]["convert"] = {"input": conv_input, "at": now_iso()}
    save_manifest(src, m)
    return "converted" + (f" (warnings: {len(warns)})" if warns else "")


# ----------------------------------------------------------------------------- ZIP


def zip_member_name(info: zipfile.ZipInfo) -> str:
    name = info.filename
    if not info.flag_bits & 0x800:
        try:
            name = name.encode("cp437").decode("utf-8")
        except (UnicodeEncodeError, UnicodeDecodeError):
            pass
    return unicodedata.normalize("NFC", name)


def cmd_unzip(args, cfg):
    done = 0
    for root_name in cfg["sources"]["raw"]:
        base = ROOT / root_name
        for z in sorted(base.rglob("*.zip")):
            if any(part.startswith(".") for part in z.relative_to(ROOT).parts):
                continue
            target = z.parent / z.stem
            zsha = sha256_file(z)
            mpath = MANIFEST_DIR / f"{slug(rel(z))}.json"
            mz = json.loads(mpath.read_text()) if mpath.exists() else {}
            if mz.get("sha256") == zsha and target.exists() and not args.force:
                continue
            with zipfile.ZipFile(z) as zf:
                infos = [i for i in zf.infolist() if not i.is_dir() and "__MACOSX" not in i.filename]
                names = [zip_member_name(i) for i in infos]
                tops = {n.split("/")[0] for n in names}
                strip = len(tops) == 1 and all("/" in n for n in names)
                for info, name in zip(infos, names):
                    relname = name.split("/", 1)[1] if strip else name
                    dest = (target / relname).resolve()
                    if not dest.is_relative_to(target.resolve()):
                        die(f"ZIP with invalid path: {name}")
                    dest.parent.mkdir(parents=True, exist_ok=True)
                    data = zf.read(info)
                    if not dest.exists() or dest.read_bytes() != data:
                        dest.write_bytes(data)
            mpath.parent.mkdir(parents=True, exist_ok=True)
            write_if_changed(mpath, json.dumps({"source": rel(z), "kind": "zip", "sha256": zsha, "extracted_to": rel(target), "at": now_iso()}, indent=2, ensure_ascii=False, sort_keys=True) + "\n")
            print(f"unzipped: {rel(z)} -> {rel(target)}/")
            done += 1
    if not done and not getattr(args, "quiet", False):
        print("unzip: nothing to do")


# ----------------------------------------------------------------------------- deterministic extraction (triples)


def iri(path: str) -> str:
    return NS_ID + path


def expand(curie: str) -> str:
    if curie.startswith("http"):
        return curie
    p, _, local = curie.partition(":")
    if p not in PREFIXES:
        die(f"unknown prefix in {curie}")
    return PREFIXES[p] + local


class G:
    """Simple triple accumulator -> deterministic Turtle (via rdflib)."""

    def __init__(self):
        import rdflib

        self.rdflib = rdflib
        self.g = rdflib.Graph()
        for k, v in PREFIXES.items():
            if k != "sh":
                self.g.bind(k, v)

    def _o(self, o):
        R = self.rdflib
        if isinstance(o, R.term.Node):
            return o
        if isinstance(o, bool):
            return R.Literal(o)
        if isinstance(o, int):
            return R.Literal(o)
        if isinstance(o, dt.date):
            return R.Literal(o.isoformat(), datatype=R.XSD.date)
        return R.Literal(str(o))

    def U(self, s: str):
        return self.rdflib.URIRef(s)

    def add(self, s: str, p: str, o, lang: str | None = None):
        if o is None or o == "":
            return
        R = self.rdflib
        obj = R.Literal(str(o), lang=lang) if lang else self._o(o)
        self.g.add((R.URIRef(s), R.URIRef(expand(p)), obj))

    def link(self, s: str, p: str, o: str):
        self.g.add((self.rdflib.URIRef(s), self.rdflib.URIRef(expand(p)), self.rdflib.URIRef(o)))

    def typ(self, s: str, cls: str):
        self.link(s, "rdf:type", expand(cls))

    def ttl(self, header: str) -> str:
        return header + self.g.serialize(format="turtle")

    def __len__(self):
        return len(self.g)


def parse_md_blocks(body: str):
    """Iterate markdown blocks: ('heading', level, text, page) | ('table', rows, page) | ('line', text, page)."""
    page = None
    lines = body.splitlines()
    i = 0
    while i < len(lines):
        ln = lines[i]
        m = re.match(r"<!-- page: (\d+) -->", ln)
        if m:
            page = int(m[1])
            i += 1
            continue
        h = re.match(r"^(#{1,6})\s+(.*\S)\s*$", ln)
        if h:
            yield ("heading", len(h[1]), h[2].strip("* "), page)
            i += 1
            continue
        if ln.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                if not re.match(r"^\|(\s*:?-+:?\s*\|)+\s*$", lines[i]):
                    cells = [c.strip().replace("\\|", "|") for c in re.split(r"(?<!\\)\|", lines[i].strip())[1:-1]]
                    rows.append(cells)
                i += 1
            yield ("table", rows, page)
            continue
        im = IMG_MARKER_RE.match(ln)
        if im and im["status"] == "DESCRITA":
            desc = []
            j = i + 1
            while j < len(lines) and lines[j].startswith(">"):
                desc.append(lines[j].lstrip("> ").rstrip())
                j += 1
            yield ("figure", nfc(im["asset"]), "\n".join(d for d in desc if d), page)
            i = j
            continue
        if ln.strip() and not ln.startswith("<!--") and not ln.startswith(">"):
            # paragraph: join consecutive lines (PDFs split sentences across lines)
            para = [ln.strip()]
            i += 1
            while i < len(lines) and lines[i].strip() and not re.match(r"^(#{1,6}\s|\||<!--|>)", lines[i]) \
                    and not re.match(r"^\s*[-*]\s", lines[i]):
                para.append(lines[i].strip())
                i += 1
            yield ("line", "\n".join(para), page)  # keep line breaks: rules with (?m)^ match per line
            continue
        i += 1


ATTACH_WORDS = r"(?:Anexo|Annex|Attachment|Appendix|Ap[eé]ndice|Exhibit|Schedule|Adjunto)"
ANEXO_RE = re.compile(ATTACH_WORDS + r"\s+(\d+|[IVX]+)\b", re.I)
PRAZO_DUR_RE = re.compile(
    r"\b(\d{1,3})\s*(?:\([^)]{1,40}\)\s*)?("
    r"dias?\s+(?:úteis|uteis|corridos)|d[ií]as?\s+(?:h[aá]biles|naturales|corridos)|(?:business|working|calendar)\s+days?"
    r"|dias?|d[ií]as|days?|horas?|hours?|meses|months?|semanas?|weeks?|anos|años|years?)\b", re.I)
PRAZO_KW_RE = re.compile(r"\b(prazos?|até|antecedência|vigência|validade|vencimento|limite|entrega|no máximo|no mínimo|dentro de|após"
                         r"|deadlines?|due|within|until|no later than|at least|at most|no less than|not less than|valid|validity|term|minimum"
                         r"|plazos?|hasta|vigencia|validez|v[aá]lid[ao]s?|como m[ií]nimo|dentro de)\b", re.I)
REQ_ID_RE = re.compile(r"\b(?:RF|RNF|REQ|RQ|RS|RSI|RN)[-_ ]?\d{1,4}(?:\.\d+)?\b")
ROMAN = {"I": 1, "II": 2, "III": 3, "IV": 4, "V": 5, "VI": 6, "VII": 7, "VIII": 8, "IX": 9, "X": 10}
GLOSS_HDR = re.compile(r"^(sigla|termo|abreviatura|acr[oô]nimo|termo/sigla|term|acronym|abbreviation|glossary|t[eé]rmino)s?$", re.I)
DEF_HDR = re.compile(r"^(descri[cç][aã]o|defini[cç][aã]o|significado|description|definition|meaning|descripci[oó]n|definici[oó]n)s?$", re.I)
MARCO_LABEL_HDR = re.compile(r"(etapa|evento|atividade|marco|fase|descri|stage|event|activity|milestone|phase|step|actividad|hito)", re.I)
MARCO_DATE_HDR = re.compile(r"(data|prazo|date|deadline|due|fecha|plazo)", re.I)


def anexo_index(sources: list[Source]) -> dict[int, str]:
    out = {}
    for s in sources:
        if s.layer != "raw":
            continue
        m = re.match(ATTACH_WORDS + r"\s+(\d+)\b", s.path.name, re.I)
        if m:
            out.setdefault(int(m[1]), s.doc_iri())
    return out


def fold(s: str) -> str:
    return unicodedata.normalize("NFKD", s).encode("ascii", "ignore").decode().lower()


def anexo_targets(entry: dict) -> list[str]:
    """Document IRIs of a configured attachment's file (or folder)."""
    arq = entry.get("file")
    if not arq:
        return []
    p = ROOT / unicodedata.normalize("NFC", arq)
    if p.is_dir():
        return sorted({NS_ID + "doc/" + slug(rel(f)) for f in walk_files(p) if f.suffix.lower() in BINARY_EXTS})
    return [NS_ID + "doc/" + slug(unicodedata.normalize("NFC", arq))]


def resolve_anexo(n: int, after: str, src: Source, cfg: dict, by_filename: dict[int, str]) -> tuple[list[str], str | None]:
    """Resolve "Anexo N" -> (documents, note). Order: name/alias right after the mention (fixes wrong
    numbering in the text) > official numbering of the document family (kb/config.yaml `attachments`) > file name."""
    entries = cfg.get("attachments") or []
    win = fold(after)
    by_alias = next((e for e in entries if any(fold(al) in win for al in e.get("aliases") or [])), None)
    fam_nums = {e["id"]: (e.get("numbers") or {}).get(src.family) for e in entries}
    by_number = next((e for e in entries if n in (fam_nums[e["id"]] or [])), None)
    family_has_list = any(fam_nums.values())
    if by_alias:
        tg = anexo_targets(by_alias)
        nota = None
        if by_number and by_number is not by_alias:
            nota = f"Anexo {n}: numbering mismatch — the text cites {by_alias['title']}, which is not Anexo {n} in the official list"
        elif not tg:
            nota = f"Anexo {n}: {by_alias['title']} (missing from raw/)"
        return tg, nota
    if by_number:
        tg = anexo_targets(by_number)
        return tg, (None if tg else f"Anexo {n}: {by_number['title']} (missing from raw/)")
    if not family_has_list and n in by_filename:
        return [by_filename[n]], None
    return [], f"Anexo {n} (unresolved)"


def section_class_for(src: Source, cfg: dict) -> str | None:
    import fnmatch

    for pat, cls in (cfg.get("section_class") or {}).items():
        if fnmatch.fnmatch(src.rel, pat) or fnmatch.fnmatch(src.path.name, pat):
            return cls
    return None


def doc_meta(src: Source, cfg: dict) -> dict:
    """Precedence metadata of the source (layer, date, version, family, overrides, inferred header)."""
    meta = {"layer": src.layer, "family": src.family, "title": src.path.stem, "date": None, "version": None,
            "version_key": "", "overrides": [], "override_refs": [], "header_inferred": False, "author": None,
            "review_every": None, "warnings": []}
    if not src.md.exists():
        return meta
    fm, _, _ = parse_frontmatter(src.md.read_text(encoding="utf-8"))
    if src.layer == "raw":
        meta["title"] = fm.get("title") or src.path.stem
        meta["date"] = (to_date(fm.get("document_date")) or None)
        meta["version"] = str(fm["version"]) if fm.get("version") not in (None, "") else None
        meta["version_key"] = version_key(meta["version"])
        meta["warnings"] = [w for w in (fm.get("warnings") or [])
                            if not (meta["version"] and str(w).startswith(("version not found", "versão não encontrada", "ambiguous version", "versão ambígua")))]
        if src.kind == "md" and not meta["date"]:
            meta["date"] = dt.date.fromtimestamp(src.md.stat().st_mtime)
        return meta
    # wiki/ note (R2)
    meta["title"] = fm.get("title") or src.path.stem
    d = to_date(fm.get("last_reviewed"))
    author = fm.get("reviewed_by")
    if not d or not author:
        meta["header_inferred"] = True
        gd, ga = None, None
        if is_git():
            r = subprocess.run(["git", "-C", str(ROOT), "log", "-1", "--format=%aI|%an", "--", src.rel], capture_output=True, text=True)
            if r.stdout.strip():
                gd, ga = r.stdout.strip().split("|", 1)
        d = d or to_date(gd) or dt.date.fromtimestamp(src.md.stat().st_mtime)
        author = author or ga or "unknown"
        missing = [k for k in ("last_reviewed", "reviewed_by") if not fm.get(k)]
        meta["warnings"].append(f"incomplete header ({', '.join(missing)}); inferred via {'git' if gd or ga else 'mtime'}")
    meta["date"] = d
    meta["author"] = str(author)
    meta["review_every"] = fm.get("review_every") or cfg.get("review_every")
    meta["note_type"] = str(fm.get("type") or "").strip().lower() or None
    if meta["note_type"] and meta["note_type"] not in ("fleeting", "literature", "permanent"):
        meta["warnings"].append(f"unknown note type '{meta['note_type']}' (use fleeting | literature | permanent)")
    meta["note_source"] = str(fm.get("source") or "").strip() or None
    ovs = fm.get("overrides") or []
    if isinstance(ovs, str):
        ovs = [ovs]
    for o in ovs:
        o = unicodedata.normalize("NFC", str(o).strip())
        cand = (src.md.parent / o)
        if (ROOT / o).exists() or o.startswith(tuple(cfg["sources"]["raw"])):
            meta["overrides"].append(NS_ID + "doc/" + slug(o))
        elif cand.exists():
            meta["overrides"].append(NS_ID + "doc/" + slug(rel(cand)))
        elif o.startswith("http"):
            meta["overrides"].append(o)
        elif o.startswith("id:"):
            meta["overrides"].append(NS_ID + o[3:])
        else:
            meta["override_refs"].append(o)
    meta["version_key"] = ""
    return meta


def det_input_hash(src: Source, cfg: dict, mapping_sha: str, meta: dict) -> str:
    parts = [EXTRACTOR_VERSION, src.family, mapping_sha,
             json.dumps({k: cfg.get(k) for k in ("section_class", "attachments", "topics", "facts")}, sort_keys=True, default=str)]
    if src.md.exists():
        parts += list(md_hashes(src.md))
    if src.binary_present and src.kind == "xlsx":
        parts.append(sha256_file(src.path))
    if src.layer == "wiki":
        parts.append(json.dumps({k: str(v) for k, v in meta.items()}, sort_keys=True))
        parts.append(",".join(sorted(source_index(cfg)[1])))  # [[links]] resolve against the current note set
    return sha256_text("|".join(parts))


def find_mapping(src: Source) -> tuple[dict | None, str]:
    if not MAPPINGS_DIR.exists():
        return None, ""
    for p in sorted(MAPPINGS_DIR.glob("*.yaml")):
        text = p.read_text(encoding="utf-8")
        m = yaml.safe_load(text) or {}
        if unicodedata.normalize("NFC", m.get("source", "")) == src.rel:
            return m, sha256_text(text)
    return None, ""


def tmpl(t: str, rec: dict, slugify: bool) -> str:
    def rep(m):
        v = cell_str(rec.get(m.group(1)))
        return slug(v, keep_dots=True) if slugify else v
    return re.sub(r"\{([^}]+)\}", rep, t)


def extract_xlsx(g: G, src: Source, doc: str, mapping: dict, cfg: dict) -> list[str]:
    import openpyxl
    from openpyxl.utils import get_column_letter

    warns = []
    if not src.binary_present:
        return ["binary missing: spreadsheet triples kept from the last extraction are not regenerated"]
    wb = openpyxl.load_workbook(str(src.path), read_only=True, data_only=True)
    for sm in mapping.get("sheets", []):
        if sm["sheet"] not in wb.sheetnames:
            warns.append(f"sheet not found: {sm['sheet']}")
            continue
        ws = wb[sm["sheet"]]
        rows = xlsx_rows(ws, cfg["xlsx"]["max_empty_rows"])
        hdr_row = sm.get("header_row") or detect_header(rows)
        hdr = next((v for i, v in rows if i == hdr_row), None)
        if hdr is None:
            warns.append(f"header not found: {sm['sheet']} row {hdr_row}")
            continue
        headers = [cell_str(h) for h in hdr]
        sec = iri(f"section/{src.family}/aba-{slug(sm['sheet'])}")
        g.typ(sec, "kb:Section")
        g.add(sec, "rdfs:label", f"Sheet {sm['sheet']}")
        g.link(sec, "kb:definedIn", doc)
        g.add(sec, "kb:sheet", sm["sheet"])
        filldown = {c: None for c in sm.get("fill_down", [])}
        count = 0
        matrix = sm.get("matrix")
        mcols = []
        if matrix:
            if "columns" in matrix:
                mcols = matrix["columns"]
            else:
                start = headers.index(matrix["from"])
                end = headers.index(matrix["to"]) + 1 if matrix.get("to") else len(headers)
                mcols = [h for h in headers[start:end] if h]
        for rownum, vals in rows:
            if rownum <= hdr_row:
                continue
            rec = {h: v for h, v in zip(headers, vals) if h}
            rec.update({f"@{get_column_letter(i + 1)}": v for i, v in enumerate(vals)})  # column by letter: {@B}
            for c in filldown:
                if cell_str(rec.get(c)):
                    filldown[c] = rec[c]
                else:
                    rec[c] = filldown[c]
            if any(not cell_str(rec.get(c)) for c in sm.get("require", [])):
                continue
            s = iri(tmpl(sm["iri"], rec, True))
            classes = sm["class"] if isinstance(sm["class"], list) else [sm["class"]]
            for c in classes:
                g.typ(s, c)
            tbs = sm.get("type_by") or []
            for tb in (tbs if isinstance(tbs, list) else [tbs]):
                cls = tb["map"].get(cell_str(rec.get(tb["column"])))
                if cls:
                    g.typ(s, cls)
            if sm.get("label"):
                g.add(s, "rdfs:label", tmpl(sm["label"], rec, False))
            g.link(s, "kb:definedIn", doc)
            g.link(s, "kb:inSection", sec)
            g.add(s, "kb:sheet", sm["sheet"])
            g.add(s, "kb:row", rownum)
            for col, spec in (sm.get("columns") or {}).items():
                v = rec.get(col)
                if not cell_str(v) or cell_str(v) in ("-", "—"):
                    continue
                spec = {"prop": spec} if isinstance(spec, str) else spec
                kind = spec.get("type", "str")
                if kind == "iri":
                    o = iri(tmpl(spec["iri"], rec, True))
                    g.link(s, spec["prop"], o)
                    if spec.get("class"):
                        g.typ(o, spec["class"])
                        g.add(o, "rdfs:label", cell_str(v))
                        g.link(o, "kb:definedIn", doc)
                        g.add(o, "kb:sheet", sm["sheet"])
                elif kind == "date":
                    d = to_date(v)
                    g.add(s, spec["prop"], d if d else cell_str(v))
                elif kind == "int":
                    try:
                        g.add(s, spec["prop"], int(float(v)))
                    except (TypeError, ValueError):
                        g.add(s, spec["prop"], cell_str(v))
                else:
                    g.add(s, spec["prop"], cell_str(v))
            for col in mcols:
                code = re.sub(r"[.;]", ",", re.sub(r"\s+", "", cell_str(rec.get(col)).upper()))
                if not code or code in ("-", "N/A", "NA"):
                    continue
                papel = iri(f"role/{slug(col)}")
                g.typ(papel, "kb:Role")
                g.add(papel, "rdfs:label", col.strip())
                g.link(papel, "kb:definedIn", doc)
                g.add(papel, "kb:sheet", sm["sheet"])
                att = iri(f"raci/{tmpl(sm['iri'], rec, True).replace('/', '-')}/{slug(col)}")
                g.typ(att, "kb:RACIAssignment")
                g.link(att, "kb:activity", s)
                g.link(att, "kb:role", papel)
                g.add(att, "kb:raciCode", code)
                g.link(att, "kb:definedIn", doc)
                g.add(att, "kb:sheet", sm["sheet"])
                g.add(att, "kb:row", rownum)
                for letter in re.findall(r"[RACI]", code):
                    prop = (matrix.get("codes") or {}).get(letter) or {"R": "kb:responsible", "A": "kb:accountable", "C": "kb:consulted", "I": "kb:informed"}[letter]
                    g.link(s, prop, papel)
            count += 1
        g.add(sec, "kb:rowCount", count)
    wb.close()
    return warns


def extract_markdown(g: G, src: Source, doc: str, body: str, cfg: dict, anexos: dict[int, str], example: bool = False):
    """Structure (sections), glossary, milestones, deadlines and cross-references from the .md."""
    sec_cls = section_class_for(src, cfg)
    is_note = src.layer == "wiki"
    base = f"section/{'note-' + src.id if is_note else src.family}"
    stack: list[tuple[int, str]] = []
    facts_acc: dict = {}
    current = None
    seen: Counter = Counter()
    sec_text: dict[str, list[str]] = defaultdict(list)
    prazo_n = 0
    for blk in parse_md_blocks(body):
        kind = blk[0]
        page = blk[-1]
        if kind == "heading":
            _, level, text, _ = blk
            num = NUMBERED_HEADING_RE.match(text)
            key = num.group(1) if num else slug(text)[:60]
            seen[key] += 1
            if seen[key] > 1:
                key = f"{key}-{seen[key]}"
            s = iri(f"{base}/{key}")
            g.typ(s, "kb:Section")
            if sec_cls:
                g.typ(s, sec_cls)
            g.add(s, "rdfs:label", text)
            if num:
                g.add(s, "kb:number", num.group(1))
            g.link(s, "kb:definedIn", doc)
            if page:
                g.add(s, "kb:page", page)
            else:
                g.add(s, "kb:order", sum(seen.values()))  # no pages (DOCX/MD): heading position
            while stack and stack[-1][0] >= level:
                stack.pop()
            if stack:
                g.link(s, "kb:parentSection", stack[-1][1])
            stack.append((level, s))
            current = s
            if not is_note and cfg.get("facts"):
                extract_facts(g, text, page, doc, cfg, facts_acc)
            continue
        if kind == "figure":  # described image/diagram -> searchable entity
            _, asset, text, _ = blk
            if text:
                fi = iri(f"figure/{src.family}/{slug(Path(asset).stem)}")
                first = re.sub(r"^\[(?:Figura|Figure|Figure)[^\]]*\]\s*", "", text.splitlines()[0]).strip()
                g.typ(fi, "kb:Figure")
                g.add(fi, "rdfs:label", first[:200] or Path(asset).stem)
                g.add(fi, "kb:excerpt", text[:2000])
                g.add(fi, "kb:originalId", asset)
                g.link(fi, "kb:definedIn", doc)
                if page:
                    g.add(fi, "kb:page", page)
                if current:
                    g.link(fi, "kb:inSection", current)
                if not is_note and cfg.get("facts"):
                    extract_facts(g, text, page, doc, cfg, facts_acc)
            continue
        if kind == "table":
            rows = blk[1]
            if len(rows) < 2:
                continue
            hdr = [c.strip() for c in rows[0]]
            gi = next((i for i, h in enumerate(hdr) if GLOSS_HDR.match(h)), None)
            di = next((i for i, h in enumerate(hdr) if DEF_HDR.match(h)), None)
            if gi is not None and di is not None:
                for r in rows[1:]:
                    if len(r) > max(gi, di) and r[gi] and r[di]:
                        t = iri(f"term/{slug(r[gi])}")
                        g.typ(t, "kb:Term")
                        g.add(t, "skos:prefLabel", r[gi], lang=(cfg.get("language") or "en")[:2])
                        g.add(t, "skos:definition", r[di], lang=(cfg.get("language") or "en")[:2])
                        g.link(t, "kb:definedIn", doc)
                        if page:
                            g.add(t, "kb:page", page)
                continue
            li = next((i for i, h in enumerate(hdr) if MARCO_LABEL_HDR.search(h)), None)
            dcol = next((i for i, h in enumerate(hdr) if MARCO_DATE_HDR.search(h) and i != li), None)
            if li is not None and dcol is not None:
                for r in rows[1:]:
                    if len(r) <= max(li, dcol) or not r[li]:
                        continue
                    ds = find_dates(r[dcol])
                    if not ds:
                        continue
                    mk = iri(f"milestone/{slug(r[li])[:80]}")
                    g.typ(mk, "kb:Milestone")
                    g.add(mk, "rdfs:label", r[li])
                    g.add(mk, "kb:date", ds[0][0])
                    g.add(mk, "kb:excerpt", " | ".join(c for c in r if c))
                    g.link(mk, "kb:definedIn", doc)
                    if page:
                        g.add(mk, "kb:page", page)
                    if current:
                        g.link(mk, "kb:inSection", current)
                    for tema, tc in (cfg.get("topics") or {}).items():
                        if re.search(tc["pattern"], fold(" ".join(c for c in r if c))):
                            ts = iri(f"topic/{slug(tema)}")
                            g.typ(ts, "kb:Deadline")
                            g.add(ts, "rdfs:label", tc.get("label") or tema)
                            g.add(ts, "kb:topic", slug(tema))
                            g.add(ts, "kb:date", ds[0][0])
                            g.add(ts, "kb:excerpt", " | ".join(c for c in r if c))
                            g.link(ts, "kb:definedIn", doc)
                            g.add(ts, "kb:originalId", f"{src.rel}#{slug(tema)}")
                            if page:
                                g.add(ts, "kb:page", page)
                            g.link(mk, "kb:aboutTopic", ts)
            text = " ".join(" ".join(r) for r in rows)
        else:
            text = blk[1]
            if is_note:
                fact = re.match(NOTE_FACT_RE, text, re.I)
                if fact and example:
                    continue  # example notes (header `example: true`) show conventions without feeding the graph
                if fact:
                    extract_note_fact(g, src, doc, fact.group(1), fact.group(2), current)
                    continue
        if current:
            sec_text[current].append(text)
        if not is_note and cfg.get("facts"):
            extract_facts(g, text, page, doc, cfg, facts_acc)
        subj = current or doc
        for m in ANEXO_RE.finditer(text):
            a = m.group(1)
            n = int(a) if a.isdigit() else ROMAN.get(a.upper(), 0)
            tgts, nota = resolve_anexo(n, text[m.end(): m.end() + 120], src, cfg, anexos)
            for tgt in tgts:
                if tgt != doc:
                    g.link(subj, "kb:references", tgt)
            if nota:
                g.add(subj, "kb:referenceText", nota)
        for rid in set(REQ_ID_RE.findall(text)):
            g.add(current or doc, "kb:mentionsId", rid)
        by_path, by_key = source_index(cfg)
        for link in re.findall(r"\]\(([^)]+\.(?:md|pdf|docx|xlsx))\)", text):
            p = (src.md.parent / unicodedata.normalize("NFC", link.replace("%20", " "))).resolve()
            tgt = by_path.get(rel(p)) if p.exists() and p.is_relative_to(ROOT) else None
            if tgt and tgt != doc:
                g.link(doc if is_note else (current or doc), "kb:linksTo" if is_note else "kb:references", tgt)
        for wl in WIKILINK_RE.findall(text):  # Zettelkasten [[links]] between notes
            key = slug(wl.strip())
            tgt = by_key.get(key) or by_key.get(re.sub(r"^\d{4}-\d{2}-\d{2}-", "", key))
            if tgt and tgt != doc:
                g.link(doc, "kb:linksTo", tgt)
            elif not tgt:
                g.add(doc, "kb:referenceText", f"[[{wl.strip()}]] (not found)")
        if kind == "line" and not is_note:
            for sent in re.split(r"(?<=[.;!?])\s+(?=[A-ZÀ-Ú0-9])", text):
                if not PRAZO_KW_RE.search(sent):
                    continue
                dates = find_dates(sent)
                durs = PRAZO_DUR_RE.findall(sent)
                if not (dates or durs):
                    continue
                prazo_n += 1
                pz = iri(f"deadline/{src.family}/p{page or 0}-{prazo_n}")
                g.typ(pz, "kb:Deadline")
                g.add(pz, "kb:excerpt", sent[:500])
                for d, _ in dates[:3]:
                    g.add(pz, "kb:date", d)
                for n, unit in durs[:3]:
                    g.add(pz, "kb:duration", f"{n} {unit}")
                g.link(pz, "kb:definedIn", doc)
                if page:
                    g.add(pz, "kb:page", page)
                if current:
                    g.link(pz, "kb:inSection", current)
                # topic: the same condition in different documents converges on one subject -> precedence/conflict
                for tema, tc in (cfg.get("topics") or {}).items():
                    if re.search(tc["pattern"], fold(sent)):
                        ts = iri(f"topic/{slug(tema)}")
                        g.typ(ts, "kb:Deadline")
                        g.add(ts, "rdfs:label", tc.get("label") or tema)
                        g.add(ts, "kb:topic", slug(tema))
                        g.add(ts, "kb:excerpt", sent[:500])
                        for d, _ in dates[:3]:
                            g.add(ts, "kb:date", d)
                        for n, unit in durs[:3]:
                            g.add(ts, "kb:duration", f"{n} {unit.lower()}")
                        g.link(ts, "kb:definedIn", doc)
                        g.add(ts, "kb:originalId", f"{src.rel}#{slug(tema)}")  # locator also for sources without pages
                        if page:
                            g.add(ts, "kb:page", page)
                        g.link(pz, "kb:aboutTopic", ts)
    if facts_acc:
        flush_facts(g, doc, cfg, facts_acc)
    for s, texts in sec_text.items():
        t = re.sub(r"\s+", " ", " ".join(texts)).strip()
        if t:
            g.add(s, "kb:excerpt", t[:1500])


NUM_WORDS = {  # pt / es / en number words (accent-free)
    "um": 1, "uma": 1, "dois": 2, "duas": 2, "tres": 3, "quatro": 4, "cinco": 5, "seis": 6, "sete": 7, "oito": 8,
    "nove": 9, "dez": 10, "doze": 12, "vinte": 20, "trinta": 30, "quarenta": 40, "cinquenta": 50, "sessenta": 60,
    "noventa": 90,
    "uno": 1, "dos": 2, "cuatro": 4, "siete": 7, "ocho": 8, "nueve": 9, "diez": 10, "veinte": 20, "treinta": 30,
    "cuarenta": 40, "cincuenta": 50, "sesenta": 60,
    "one": 1, "two": 2, "three": 3, "four": 4, "five": 5, "six": 6, "seven": 7, "eight": 8, "nine": 9, "ten": 10,
    "twelve": 12, "twenty": 20, "thirty": 30, "forty": 40, "fifty": 50, "sixty": 60, "ninety": 90}


def fold_keep(s: str) -> str:
    """fold() preserving length (indices in the folded text are valid in the original, to quote the literal excerpt)."""
    return "".join("-" if c in "–—‐‑" else (fold(c) or " ")[0] for c in s)


def fact_items(v: str) -> list[str]:
    """Split an enumeration ("a, b and c" / "a, b e c" / "a, b y c") into items."""
    v = re.sub(r"\b(?:um|uma|an?|un|una)\s+(?:ambiente|environment|entorno)\s+(?:de|of|for)\s+", "", v)
    v = re.sub(r"\([^)]*\)", "", v)
    seps = r",|;|\be\b|\bou\b|\band\b|\bor\b|\by\b"
    if (load_config().get("language") or "")[:2] == "es":
        seps += r"|\bo\b"  # Spanish "o" (or) only in Spanish workspaces: in Portuguese it is the article "the"
    return [x.strip(" -") for x in re.split(seps, v) if x.strip(" -")]


def fact_value(v: str, fc: dict) -> str | None:
    """Normalize the captured value: count (list items), months (years -> months), number (spelled out -> digits)."""
    v = re.sub(r"\s+", " ", v).strip()
    norm = fc.get("normalize")
    if norm == "count":
        n = len(fact_items(v))
        return str(n) if n >= fc.get("min", 2) else None  # a 1-item list is not an enumeration
    if norm == "acronyms":
        sig = [x.upper() for x in re.findall(r"\(([a-z]{2,6})\)", v)]
        extra = [w for w in (fc.get("extras") or []) if w in v]
        return ", ".join(sig) + (" + " + " + ".join(extra) if extra else "") if sig else None
    if norm in ("months", "number"):
        m = re.search(r"(\d+|" + "|".join(sorted(NUM_WORDS, key=len, reverse=True)) + r")\s*(?:\([a-z ]+\)\s*)?"
                      r"(meses|mes|months?|anos|ano|years?)?", v)
        if not m:
            return None
        n = int(m.group(1)) if m.group(1).isdigit() else NUM_WORDS[m.group(1)]
        if norm == "months" and (m.group(2) or "").startswith(("ano", "year")):
            n *= 12
        return str(n)
    return v[:160]


def extract_facts(g: G, text: str, page: int | None, doc: str, cfg: dict, acc: dict):
    """Quantitative facts (kb/config.yaml `facts`): concept + value, searched in a window around the context.
    The same concept in different documents converges on id:topic/<fact>, so divergent values
    (e.g. 4 × 3 environments) become a precedence conflict; two values in the same document show up in
    kb/queries/inconsistencias-internas.rq."""
    fs = fold_keep(text)
    for nome, fc in cfg["facts"].items():
        for pm in re.finditer(fc["context"], fs):
            ws, we = max(0, pm.start() - fc.get("before", 50)), min(len(fs), pm.end() + fc.get("after", 250))
            vm = re.search(fc["value"], fs[ws:we])
            if not vm:
                continue
            trecho = text[ws:we].strip()
            if fc.get("normalize") == "count_distinct":
                a = acc.setdefault(nome, {"items": set(), "pages": set(), "trechos": []})
                a["items"].update(i for i in fact_items(vm.group(1)))
                if page:
                    a["pages"].add(page)
                a["trechos"].append(trecho)
                continue
            val = fact_value(vm.group(1), fc)
            if val is None:
                continue
            emit_fact(g, nome, fc, val, trecho, page, doc)


def emit_fact(g: G, nome: str, fc: dict, val: str, trecho: str, page, doc: str):
    ts = iri(f"topic/{slug(nome)}")
    g.typ(ts, "kb:Fact")
    g.add(ts, "rdfs:label", fc.get("label") or nome)
    g.add(ts, "kb:topic", slug(nome))
    g.add(ts, "kb:value", f"{val} {fc.get('unit', '')}".strip())
    g.add(ts, "kb:excerpt", trecho[:500])
    g.link(ts, "kb:definedIn", doc)
    g.add(ts, "kb:originalId", f"{doc.rsplit('/', 1)[-1]}#{slug(nome)}")  # locator also for sources without pages
    for p in (page if isinstance(page, set) else {page}):
        if p:
            g.add(ts, "kb:page", p)


def flush_facts(g: G, doc: str, cfg: dict, acc: dict):
    for nome, a in acc.items():
        emit_fact(g, nome, cfg["facts"][nome], str(len(a["items"])), " … ".join(a["trechos"][:3]), a["pages"], doc)


NOTE_FACT_KINDS = {  # note line conventions (en / pt / es) -> canonical kind
    "deadline": "deadline", "prazo": "deadline", "plazo": "deadline",
    "decision": "decision", "decisao": "decision",
    "hypothesis": "hypothesis", "hipotese": "hypothesis", "hipotesis": "hypothesis",
    "assumption": "assumption", "premissa": "assumption", "supuesto": "assumption",
    "term": "term", "termo": "term", "termino": "term",
    "condition": "condition", "condicao": "condition", "condicion": "condition",
}
NOTE_FACT_RE = r"^[-*]\s*(" + "|".join(sorted({k for k in NOTE_FACT_KINDS}, key=len, reverse=True)).replace("c", "[cç]").replace("a", "[aã]").replace("o", "[oóõ]").replace("e", "[eé]").replace("i", "[ií]") + r")\s*:\s*(.+)$"


_SRC_INDEX: tuple[dict, dict] | None = None


def source_index(cfg: dict) -> tuple[dict, dict]:
    """(relative path of binary or .md -> source IRI, note key -> note IRI) for link resolution."""
    global _SRC_INDEX
    if _SRC_INDEX is None:
        sources, _, _ = discover(cfg)
        by_path, by_key = {}, {}
        for s in sources:
            by_path[s.rel] = s.doc_iri()
            if s.md.exists():
                by_path[rel(s.md)] = s.doc_iri()
            if s.layer == "wiki":
                stem = slug(s.path.stem)
                by_key[stem] = s.doc_iri()
                by_key.setdefault(re.sub(r"^\d{4}-\d{2}-\d{2}-", "", stem), s.doc_iri())
                fm, _, _ = parse_frontmatter(s.md.read_text(encoding="utf-8", errors="replace"))
                if fm.get("title"):
                    by_key.setdefault(slug(str(fm["title"])), s.doc_iri())
        _SRC_INDEX = (by_path, by_key)
    return _SRC_INDEX


WIKILINK_RE = re.compile(r"\[\[([^\]|#]+)(?:#[^\]|]*)?(?:\|[^\]]*)?\]\]")


def extract_note_fact(g: G, src: Source, doc: str, kind: str, rest: str, current: str | None):
    k = NOTE_FACT_KINDS.get(slug(kind).replace("-", ""), "hypothesis")
    lang = (load_config().get("language") or "en")[:2]
    if k == "deadline" and "=" in rest:
        label, val = [x.strip() for x in rest.split("=", 1)]
        s = iri(f"milestone/{slug(label)[:80]}")
        g.typ(s, "kb:Milestone")
        g.add(s, "rdfs:label", label)
        d = to_date(val)
        g.add(s, "kb:date", d if d else val)
    elif k == "condition" and "=" in rest:
        label, val = [x.strip() for x in rest.split("=", 1)]
        s = iri(f"topic/{slug(label)}")
        g.typ(s, "kb:Fact")
        g.add(s, "kb:topic", slug(label))
        g.add(s, "kb:value", val)
    elif k == "term" and "=" in rest:
        label, val = [x.strip() for x in rest.split("=", 1)]
        s = iri(f"term/{slug(label)}")
        g.typ(s, "kb:Term")
        g.add(s, "skos:prefLabel", label, lang=lang)
        g.add(s, "skos:definition", val, lang=lang)
    else:
        cls = {"decision": "kb:Decision", "hypothesis": "kb:Hypothesis", "assumption": "kb:Hypothesis"}.get(k, "kb:Hypothesis")
        s = iri(f"{k}/{src.id}/{sha256_text(rest)[:10]}")
        g.typ(s, cls)
        g.add(s, "rdfs:label", rest[:200])
        g.add(s, "kb:excerpt", rest)
        if k == "assumption":
            g.add(s, "kb:hypothesisKind", "assumption")
    g.link(s, "kb:definedIn", doc)
    g.add(s, "kb:originalId", f"{src.rel}#{slug(rest)[:40]}")
    if current:
        g.link(s, "kb:inSection", current)


def extract_det(src: Source, cfg: dict, anexos: dict[int, str], force: bool = False) -> str:
    if not src.md.exists():
        return "no .md (run convert)"
    m = load_manifest(src)
    mapping, mapping_sha = find_mapping(src) if src.kind == "xlsx" else (None, "")
    meta = doc_meta(src, cfg)
    inp = det_input_hash(src, cfg, mapping_sha, meta)
    if not force and m["steps"].get("det", {}).get("input") == inp and src.det_ttl().exists():
        return "unchanged"
    fm, fm_raw, body = parse_frontmatter(src.md.read_text(encoding="utf-8"))
    g = G()
    doc = src.doc_iri()
    warns = list(meta["warnings"])
    if src.layer == "raw":
        g.typ(doc, "kb:Document")
        g.add(doc, "dcterms:source", src.rel)
        g.add(doc, "kb:family", src.family)
        if meta["version"]:
            g.add(doc, "kb:version", meta["version"])
    else:
        g.typ(doc, "kb:WikiNote")
        g.add(doc, "dcterms:source", src.rel)
        pessoa = iri(f"person/{slug(meta['author'])}")
        g.typ(pessoa, "prov:Agent")
        g.add(pessoa, "rdfs:label", meta["author"])
        g.link(doc, "prov:wasAttributedTo", pessoa)
        g.add(doc, "kb:reviewEvery", str(meta["review_every"]))
        if meta.get("note_type"):
            g.add(doc, "kb:noteType", meta["note_type"])
        if meta.get("note_source"):
            tgt = source_index(cfg)[0].get(unicodedata.normalize("NFC", meta["note_source"]))
            if tgt:
                g.link(doc, "kb:references", tgt)
            else:
                g.add(doc, "kb:referenceText", f"source: {meta['note_source']} (not found)")
        if meta["header_inferred"]:
            g.add(doc, "kb:headerInferred", True)
        for o in meta["overrides"]:
            g.link(doc, "kb:overrides", o)
        for o in meta["override_refs"]:
            g.add(doc, "kb:overridesRef", o)
    g.add(doc, "dcterms:title", meta["title"])
    g.add(doc, "rdfs:label", meta["title"])
    g.add(doc, "kb:layer", src.layer)
    if meta["date"]:
        g.add(doc, "dcterms:modified", meta["date"])
    if src.kind == "xlsx":
        if mapping:
            warns += extract_xlsx(g, src, doc, mapping, cfg)
        else:
            warns.append("spreadsheet without mapping in kb/mappings/ (structure only)")
        # sheets as sections (from the .md text) without descending into rows
        for blk in parse_md_blocks(body):
            if blk[0] == "heading" and blk[2].startswith("Aba: "):
                name = re.sub(r" \(oculta\)$", "", blk[2][5:])
                s = iri(f"section/{src.family}/aba-{slug(name)}")
                g.typ(s, "kb:Section")
                g.add(s, "rdfs:label", f"Sheet {name}")
                g.link(s, "kb:definedIn", doc)
                g.add(s, "kb:sheet", name)
    else:
        extract_markdown(g, src, doc, body, cfg, anexos, example=bool(fm.get("example")))
    header = f"# Generated by kb.py extract-det (v{EXTRACTOR_VERSION}) from {src.rel}. Do not edit.\n"
    write_if_changed(src.det_ttl(), g.ttl(header))
    m["meta"] = {k: (v.isoformat() if isinstance(v, dt.date) else v) for k, v in meta.items()}
    m["meta"]["warnings"] = warns
    m["steps"]["det"] = {"input": inp, "at": now_iso(), "triples": len(g)}
    body_sha, fm_sha = md_hashes(src.md)
    m["body_sha256"], m["frontmatter_sha256"] = body_sha, fm_sha
    save_manifest(src, m)
    return f"extracted ({len(g)} triples{', warnings: ' + str(len(warns)) if warns else ''})"


# ----------------------------------------------------------------------------- in-memory graph


def note_is_stale(meta: dict) -> tuple[bool, int]:
    d = to_date(meta.get("date"))
    if not d or meta.get("layer") != "wiki":
        return False, 0
    due = d + parse_duration(meta.get("review_every"), 14)
    return due < today(), (today() - due).days


def graph_meta_quads(gname: str, src_id: str, method: str, meta: dict, doc_iri: str):
    from pyoxigraph import Literal, NamedNode, Quad

    g = NamedNode(gname)
    R = lambda x: NamedNode(expand(x))  # noqa: E731
    s = NamedNode(gname)
    q = [Quad(s, R("kb:layer"), Literal(meta.get("layer", "raw")), g),
         Quad(s, R("kb:extractionMethod"), Literal(method), g),
         Quad(s, R("kb:graphSource"), NamedNode(doc_iri), g),
         Quad(s, R("kb:family"), Literal(meta.get("family") or ""), g)]
    if meta.get("date"):
        q.append(Quad(s, R("dcterms:modified"), Literal(str(meta["date"]), datatype=R("xsd:date")), g))
    if meta.get("version"):
        q.append(Quad(s, R("kb:version"), Literal(str(meta["version"])), g))
    q.append(Quad(s, R("kb:versionKey"), Literal(meta.get("version_key") or ""), g))
    for o in meta.get("overrides") or []:
        q.append(Quad(s, R("kb:overrides"), NamedNode(o), g))
    if meta.get("header_inferred"):
        q.append(Quad(s, R("kb:headerInferred"), Literal(True), g))
    stale, days = note_is_stale(meta)
    if stale:
        q.append(Quad(s, R("kb:reviewOverdue"), Literal(True), g))
        q.append(Quad(s, R("kb:daysOverdue"), Literal(days), g))
    return q


def triple_files() -> list[tuple[str, str, Path]]:
    """(method, source_id, path) of every triple .ttl."""
    out = []
    for method in ("det", "llm"):
        d = TRIPLES_DIR / method
        if d.exists():
            out += [(method, p.stem, p) for p in sorted(d.glob("*.ttl"))]
    return out


def all_manifests() -> dict[str, dict]:
    out = {}
    if MANIFEST_DIR.exists():
        for p in MANIFEST_DIR.glob("*.json"):
            m = json.loads(p.read_text(encoding="utf-8"))
            if m.get("kind") != "zip":
                out[m["id"]] = m
    return out


def build_store(include_ontology: bool = True):
    from pyoxigraph import NamedNode, RdfFormat, Store

    store = Store()
    if include_ontology:
        for op in ontology_files():
            store.load(op.read_bytes(), format=RdfFormat.TURTLE, to_graph=NamedNode("urn:kb:ontology"))
    manifests = all_manifests()
    graphs = {}
    for method, sid, p in triple_files():
        gname = f"urn:kb:{method}:{sid}"
        try:
            store.load(p.read_bytes(), format=RdfFormat.TURTLE, to_graph=NamedNode(gname))
        except SyntaxError as e:
            print(f"WARNING: {rel(p)} skipped (invalid Turtle: {e}). Run kb.py validate.", file=sys.stderr)
            continue
        m = manifests.get(sid, {})
        meta = m.get("meta", {"layer": m.get("layer", "raw")})
        layer = m.get("layer", "raw")
        doc_iri = NS_ID + ("note/" if layer == "wiki" else "doc/") + sid
        for q in graph_meta_quads(gname, sid, method, meta, doc_iri):
            store.add(q)
        graphs[gname] = {**meta, "source": m.get("source", sid), "method": method, "doc": doc_iri}
    return store, graphs


def term_str(t) -> str:
    if t is None:
        return ""
    v = getattr(t, "value", str(t))
    return v


def short(v: str) -> str:
    for k, ns in PREFIXES.items():
        if v.startswith(ns):
            return f"{k}:{v[len(ns):]}"
    return v


def run_query(store, q: str):
    if "PREFIX" not in q.upper():
        q = SPARQL_PREFIXES + q
    return store.query(q, use_default_graph_as_union=True)


# ----------------------------------------------------------------------------- precedence (mirror of kb/queries/precedencia.rq)

PROV_PROPS = {expand(p) for p in ("rdf:type", "kb:definedIn", "kb:page", "kb:row", "kb:sheet", "kb:inSection",
                                  "kb:parentSection", "prov:wasDerivedFrom", "kb:originalId")}


WIKI_HIDDEN = {expand(p) for p in ("rdf:type", "kb:sheet", "kb:row", "kb:page", "kb:definedIn")}


def beats(g2: dict, g1: dict, s: str) -> bool:
    """Does g2 beat g1 for subject s? (same rule as precedencia.rq)"""
    if g2["layer"] == "wiki" and (s in (g2.get("overrides") or []) or g1["doc"] in (g2.get("overrides") or [])):
        return True
    if g1["layer"] == "wiki" and (s in (g1.get("overrides") or []) or g2["doc"] in (g1.get("overrides") or [])):
        return False
    if g2["layer"] == "raw" and g1["layer"] == "wiki":
        return True
    if g2["layer"] == g1["layer"]:
        return (str(g2.get("date") or ""), g2.get("version_key") or "") > (str(g1.get("date") or ""), g1.get("version_key") or "")
    return False


def resolve(values_by_graph: dict[str, set], graphs: dict, s: str) -> set[str]:
    """Winning graphs among those asserting values for (s, p)."""
    gs = list(values_by_graph)
    return {g for g in gs if not any(beats(graphs[g2], graphs[g], s) for g2 in gs if g2 != g)}


# ----------------------------------------------------------------------------- commands


def fmt_table(headers: list[str], rows: list[list[str]], maxw: int = 80) -> str:
    rows = [[(c if len(c) <= maxw else c[: maxw - 1] + "…") for c in r] for r in rows]
    w = [max([len(h)] + [len(r[i]) for r in rows]) for i, h in enumerate(headers)]
    line = lambda r: " | ".join(c.ljust(w[i]) for i, c in enumerate(r))  # noqa: E731
    return "\n".join([line(headers), "-+-".join("-" * x for x in w)] + [line(r) for r in rows])


def source_state(src: Source, cfg: dict, anexos) -> dict:
    m = load_manifest(src)
    st = {"fonte": src.rel, "tipo": src.kind}
    if src.kind in CONVERTERS:
        if not src.binary_present:
            st["convert"] = "md only" if src.md.exists() else "MISSING"
        else:
            conv = m["steps"].get("convert", {})
            st["convert"] = "ok" if conv.get("input") == f"{sha256_file(src.path)}:{CONVERTER_VERSION}" and src.md.exists() else "PENDING"
    else:
        st["convert"] = "-"
    pend = 0
    if src.md.exists() and src.kind not in ("note", "md"):
        _, body = split_frontmatter(src.md.read_text(encoding="utf-8"))
        pend = sum(1 for b in parse_image_blocks(body).values() if b["status"] == "PENDENTE")
    st["imagens"] = f"{pend} pending" if pend else "ok" if src.md.exists() else "-"
    if not src.md.exists():
        st["det"] = "PENDING"
    else:
        mapping, msha = find_mapping(src) if src.kind == "xlsx" else (None, "")
        inp = det_input_hash(src, cfg, msha, doc_meta(src, cfg))
        st["det"] = "ok" if m["steps"].get("det", {}).get("input") == inp and src.det_ttl().exists() else "PENDING"
    llm = m["steps"].get("llm")
    if src.llm_ttl().exists() and llm:
        st["llm"] = "ok" if src.md.exists() and llm.get("input") == md_hashes(src.md)[0] else "OUTDATED"
    elif src.llm_ttl().exists():
        st["llm"] = "unregistered"
    else:
        st["llm"] = "-"
    return st


def cmd_status(args, cfg):
    sources, conflicts, unsupported = discover(cfg)
    anexos = anexo_index(sources)
    rows = []
    for s in sources:
        st = source_state(s, cfg, anexos)
        rows.append([st["fonte"], st["tipo"], st["convert"], st["imagens"], st["det"], st["llm"]])
    print(f"Workspace: {ROOT.name}  |  mode: {'git' if is_git() else 'no git (OneDrive)'}  |  sources: {len(sources)}")
    print(f"Engine: {'workspace copy' if IS_COPY else 'central'} ({KB_CMD})  |  engine {ENGINE_VERSION}  |  kb-setup {KB_SETUP_VERSION}"
          f"  |  workspace files: engine {cfg.get('engine_version') or '?'}, kb-setup {cfg.get('kb_setup_version') or '?'}")
    vst, vdetail = vendor_state()
    if vst != "absent":
        print(f"Engine copy: {rel(VENDOR_DIR)}/ {vst}" + (f" ({vdetail})" if vdetail else ""))
    print()
    print(fmt_table(["source", "kind", "convert", "images", "det", "llm"], rows, 70))
    alerts = []
    lk = read_lock()
    if lk:
        who = "you" if lock_is_mine(lk) else f"{lk.get('name') or lk.get('user')} ({lk.get('user')}@{lk.get('host')})"
        alerts.append(f"LOCK held by {who} since {lk.get('started')} [{lk.get('command')}]")
    for c in conflicts:
        alerts.append(f"OneDrive CONFLICT COPY: {c} (resolve manually and delete the copy)")
    by_sha = defaultdict(list)
    for s in sources:
        if s.binary_present:
            by_sha[sha256_file(s.path)].append(s.rel)
    for sha, files in by_sha.items():
        if len(files) > 1:
            alerts.append("DUPLICATES (identical content): " + " = ".join(files) + " — run extract-llm on only one of them")
    manifests = all_manifests()
    known = {s.id for s in sources}
    for sid, m in manifests.items():
        if sid not in known:
            alerts.append(f"SOURCE REMOVED: {m.get('source')} (the next update cleans its manifest/triples)")
    no_version = []
    fam_size = Counter(s.family for s in sources if s.layer == "raw")
    for s in sources:
        m = manifests.get(s.id, {})
        for w in (m.get("meta") or {}).get("warnings", []):
            if w.startswith(("version not found", "versão não encontrada")):  # legacy pt prefix: older .md frontmatter
                if fam_size[s.family] > 1:
                    no_version.append(s.path.name)
            elif not w.startswith(("incomplete header", "header incompleto")):  # reported as "INFERRED HEADER" below
                alerts.append(f"WARNING {s.rel}: {w}")
        if s.kind == "xlsx" and not find_mapping(s)[0]:
            alerts.append(f"NO MAPPING: {s.rel} (propose kb/mappings/*.yaml)")
        if s.layer == "wiki":
            meta = doc_meta(s, cfg)
            stale, days = note_is_stale({**meta, "date": meta["date"]})
            if stale:
                alerts.append(f"OVERDUE NOTE: {s.rel} ({days} days overdue; last reviewed {meta['date']} by {meta['author']})")
            if meta["header_inferred"]:
                alerts.append(f"INFERRED HEADER: {s.rel} — {'; '.join(meta['warnings'])}")
    if no_version:
        alerts.append("NO VERSION in a multi-version family (the version breaks ties between equal dates; set `version` + "
                      f"manual_fields): {', '.join(no_version)}")
    for u in unsupported:
        alerts.append(f"UNSUPPORTED (skipped): {u}")
    print("\nAlerts:" if alerts else "\nAlerts: none")
    for a in alerts:
        print(f"  - {a}")


def remove_vanished(sources: list[Source]) -> list[str]:
    known = {s.id for s in sources}
    removed = []
    for sid, m in all_manifests().items():
        if sid in known:
            continue
        for p in (TRIPLES_DIR / "det" / f"{sid}.ttl", TRIPLES_DIR / "llm" / f"{sid}.ttl", MANIFEST_DIR / f"{sid}.json"):
            if p.exists():
                p.unlink()
        removed.append(m.get("source", sid))
    return removed


def select(sources: list[Source], path: str | None) -> list[Source]:
    if not path:
        return sources
    target = unicodedata.normalize("NFC", (ROOT / path).resolve().relative_to(ROOT).as_posix()) if (ROOT / path).exists() else unicodedata.normalize("NFC", path)
    sel = [s for s in sources if s.rel == target or s.rel.startswith(target.rstrip("/") + "/") or rel(s.md) == target or s.id == path]
    if not sel:
        die(f"no source matches {path}")
    return sel


def cmd_convert(args, cfg):
    with CuratorLock(cfg, "convert"):
        sources, _, _ = discover(cfg)
        reuse = collect_descriptions(sources)
        for s in select(sources, args.path):
            if s.kind in CONVERTERS:
                r = convert_source(s, cfg, reuse, args.force)
                if r != "unchanged" or args.verbose:
                    print(f"{r:<12} {s.rel}")
        n = propagate_descriptions(sources)
        if n:
            print(f"descriptions reused from identical images: {n}")


def cmd_extract_det(args, cfg):
    with CuratorLock(cfg, "extract-det"):
        sources, _, _ = discover(cfg)
        anexos = anexo_index(sources)
        for s in select(sources, args.path):
            r = extract_det(s, cfg, anexos, args.force)
            if r != "unchanged" or args.verbose:
                print(f"{s.rel}: {r}")


def run_pipeline(args, cfg, command: str):
    with CuratorLock(cfg, command):
        cmd_unzip(argparse.Namespace(force=False, quiet=True), cfg)
        sources, conflicts, _ = discover(cfg)
        removed = remove_vanished(sources)
        for r in removed:
            print(f"removed from KB: {r}")
        reuse = collect_descriptions(sources)
        targets = select(sources, getattr(args, "path", None))
        changed = bool(removed)
        for s in targets:
            if s.kind in CONVERTERS:
                r = convert_source(s, cfg, reuse, getattr(args, "force", False))
                if r != "unchanged":
                    print(f"convert: {s.rel}: {r}")
                    changed = True
        if propagate_descriptions(sources):
            changed = True
        anexos = anexo_index(sources)
        for s in targets:
            r = extract_det(s, cfg, anexos, getattr(args, "force", False))
            if r != "unchanged":
                print(f"extract-det: {s.rel}: {r}")
                changed = True
        if not changed and (WIKI_OUT / "index.md").exists():
            print("nothing to process (no source changed)")
        else:
            ok = do_validate(cfg, quiet=True)
            print("validate: " + ("ok" if ok else "VIOLATIONS (run kb.py validate)"))
            n = do_wiki(cfg)
            print(f"wiki: {n} files changed in kb/wiki/")
        pend = []
        for s in sources:
            st = source_state(s, cfg, anexos)
            if st["imagens"].endswith(" pending"):
                pend.append(f"{s.rel}: {st['imagens']}")
            if st["llm"] == "OUTDATED":
                pend.append(f"{s.rel}: llm triples outdated (run /kb extract-llm)")
        if pend:
            print("\nPending for Claude:")
            for p in pend:
                print(f"  - {p}")
        if conflicts:
            print("\nOneDrive conflict copies: " + ", ".join(conflicts))


def cmd_update(args, cfg):
    run_pipeline(args, cfg, "update")


def cmd_ingest(args, cfg):
    run_pipeline(args, cfg, "ingest")


def cmd_lock(args, cfg):
    lk = read_lock()
    if lk and not lock_is_mine(lk):
        die(f"lock held by {lk.get('user')}@{lk.get('host')} since {lk.get('started')}", 2)
    if not lk:
        LOCK_FILE.parent.mkdir(parents=True, exist_ok=True)
        LOCK_FILE.write_text(json.dumps({**lock_identity(), "started": now_iso(), "command": args.command_name or "manual"}, ensure_ascii=False) + "\n")
    print("lock acquired" if not lk else "you already hold the lock")


def cmd_unlock(args, cfg):
    lk = read_lock()
    if not lk:
        print("no lock")
        return
    if not lock_is_mine(lk) and not args.force:
        die(f"lock held by {lk.get('user')}@{lk.get('host')} since {lk.get('started')} ({lock_age(lk)} ago); use --force only after confirming with the holder", 2)
    LOCK_FILE.unlink()
    print("lock released")


def cmd_pending_images(args, cfg):
    sources, _, _ = discover(cfg)
    rows, seen = [], {}
    for s in select(sources, args.path):
        if not s.md.exists() or s.kind in ("note", "md"):
            continue
        imgs = load_manifest(s).get("images", {})
        _, body = split_frontmatter(s.md.read_text(encoding="utf-8"))
        for asset, b in parse_image_blocks(body).items():
            if b["status"] != "PENDENTE":
                continue
            sha = imgs.get(asset, "")
            dup = seen.get(sha)
            if not dup:
                seen[sha] = f"{rel(s.md)}::{asset}"
            rows.append({"md": rel(s.md), "asset": asset, "png": rel(s.md.parent / asset), "loc": b["loc"], "sha": sha[:12], "duplicata_de": dup})
    if args.json:
        print(json.dumps(rows, ensure_ascii=False, indent=1))
        return
    uniq = [r for r in rows if not r["duplicata_de"]]
    print(f"{len(rows)} PENDENTE markers ({len(uniq)} unique images; duplicates receive the description automatically)")
    print(fmt_table(["md", "asset", "loc"], [[r["md"], r["asset"], r["loc"]] for r in uniq], 90))


def page_text(body: str, page: int) -> str:
    m = re.search(rf"<!-- page: {page} -->\n(.*?)(?=\n<!-- page: \d+ -->|\Z)", body, re.S)
    return m.group(1).strip() if m else ""


def cmd_image_context(args, cfg):
    args.asset = nfc(args.asset)
    md = ROOT / args.md
    _, body = split_frontmatter(md.read_text(encoding="utf-8"))
    blk = parse_image_blocks(body).get(args.asset)
    if not blk:
        die("marker not found")
    print(f"PNG: {rel(md.parent / args.asset)}\nlocation: {blk['loc']}  status: {blk['status']}\n")
    mp = re.match(r"page (\d+)", blk["loc"])
    if mp:
        print(page_text(body, int(mp[1])))
    else:
        i = body.find(marker(args.asset, blk["loc"], blk["status"])[:-4])
        print(body[max(0, i - 1500): i + 1500])


def cmd_page_text(args, cfg):
    _, body = split_frontmatter((ROOT / args.md).read_text(encoding="utf-8"))
    print(page_text(body, args.page))


def md_lock(md: Path):
    lp = Path(tempfile.gettempdir()) / f"kb-{sha256_text(str(md))[:16]}.lock"
    f = open(lp, "w")
    fcntl.flock(f, fcntl.LOCK_EX)
    return f


def set_image_in(md: Path, asset: str, status: str, desc_lines: list[str], note: str = "") -> bool:
    fh = md_lock(md)
    try:
        text = md.read_text(encoding="utf-8")
        fm, body = split_frontmatter(text)
        lines = body.splitlines()
        for i, ln in enumerate(lines):
            m = IMG_MARKER_RE.match(ln)
            if m and nfc(m["asset"]) == nfc(asset):
                j = i + 1
                while j < len(lines) and lines[j].startswith(">"):
                    j += 1
                lines[i:j] = [marker(asset, m["loc"], status, note)] + [
                    re.sub(r"^> \[Figura[^\]]*\]", f"> [Figura {loc_label(m['loc'])}]", d) if k == 0 else d for k, d in enumerate(desc_lines)]
                new = ("---\n" + fm + "---\n" if fm else "") + "\n".join(lines) + ("\n" if body.endswith("\n") else "")
                md.write_text(new, encoding="utf-8")
                return True
        return False
    finally:
        fcntl.flock(fh, fcntl.LOCK_UN)
        fh.close()


def cmd_set_image(args, cfg):
    args.asset = nfc(args.asset)
    md = ROOT / args.md
    text = args.text if args.text is not None else sys.stdin.read()
    text = text.strip()
    if args.ignore is not None:
        status, desc, note = "IGNORADA", [], f"({args.ignore})" if args.ignore else ""
    else:
        if not text:
            die("empty description (use --text or stdin)")
        status, note = "DESCRITA", ""
        body_lines = text.splitlines()
        _, b = split_frontmatter(md.read_text(encoding="utf-8"))
        blk = parse_image_blocks(b).get(args.asset)
        if not blk:
            die(f"marker not found: {args.asset} in {args.md}")
        if not body_lines[0].startswith("[Figura"):
            body_lines[0] = f"[Figura {loc_label(blk['loc'])}] " + body_lines[0]
        desc = ["> " + ln if ln.strip() else ">" for ln in body_lines]
    if not set_image_in(md, args.asset, status, desc, note):
        die(f"marker not found: {args.asset} in {args.md}")
    print(f"{status}: {args.md} :: {args.asset}")
    # propagate to identical images (same sha) still pending in other .md files
    sources, _, _ = discover(cfg)
    src_md = md.resolve()
    sha = None
    for s in sources:
        if s.md.resolve() == src_md:
            sha = load_manifest(s).get("images", {}).get(args.asset)
    if not sha:
        return
    for s in sources:
        if not s.md.exists() or s.kind in ("note", "md"):
            continue
        for asset, isha in load_manifest(s).get("images", {}).items():
            if isha != sha or (s.md.resolve() == src_md and asset == args.asset):
                continue
            _, b = split_frontmatter(s.md.read_text(encoding="utf-8"))
            blk = parse_image_blocks(b).get(asset)
            if blk and blk["status"] == "PENDENTE":
                set_image_in(s.md, asset, status, desc, note)
                print(f"  {status} (duplicate): {rel(s.md)} :: {asset}")


def cmd_outline(args, cfg):
    import openpyxl

    sources, _, _ = discover(cfg)
    for s in select(sources, args.path):
        print(f"\n=== {s.rel} [{s.kind}]")
        if s.kind == "xlsx" and s.binary_present:
            wb = openpyxl.load_workbook(str(s.path), read_only=True, data_only=True)
            for ws in wb.worksheets:
                rows = xlsx_rows(ws, cfg["xlsx"]["max_empty_rows"])
                hdr = detect_header(rows)
                h = next((v for i, v in rows if i == hdr), [])
                print(f"  sheet '{ws.title}': {len(rows)} non-empty rows; header row {hdr}: {[cell_str(x) for x in h if cell_str(x)]}")
            wb.close()
            continue
        if not s.md.exists():
            print("  (no .md)")
            continue
        _, body = split_frontmatter(s.md.read_text(encoding="utf-8"))
        for blk in parse_md_blocks(body):
            if blk[0] == "heading":
                print(f"  {'  ' * (blk[1] - 1)}{blk[2]}" + (f"  (p.{blk[3]})" if blk[3] else ""))


def cmd_llm_done(args, cfg):
    sources, _, _ = discover(cfg)
    sel = select(sources, args.source)
    if len(sel) != 1:
        die("specify exactly one source")
    s = sel[0]
    if not s.llm_ttl().exists():
        die(f"file not found: {rel(s.llm_ttl())}")
    m = load_manifest(s)
    m["steps"]["llm"] = {"input": md_hashes(s.md)[0], "at": now_iso()}
    save_manifest(s, m)
    print(f"llm recorded: {s.rel}")


def cmd_llm_context(args, cfg):
    sources, _, _ = discover(cfg)
    sel = select(sources, args.source)
    if len(sel) != 1:
        die("specify exactly one source")
    s = sel[0]
    print(json.dumps({"fonte": s.rel, "md": rel(s.md), "saida": rel(s.llm_ttl()), "doc_iri": s.doc_iri(),
                      "familia": s.family, "ontologies": [str(p) for p in ontology_files()], "shapes": [str(p) for p in shapes_files()],
                      "base_secao": f"{NS_ID}secao/{s.family}/"}, ensure_ascii=False, indent=1))


def do_validate(cfg, quiet=False) -> bool:
    import rdflib
    from pyshacl import validate

    data = rdflib.Graph()
    origin = defaultdict(set)
    bad = []
    for method, sid, p in triple_files():
        try:
            g = rdflib.Graph().parse(p, format="turtle")
        except Exception as e:  # rdflib BadSyntax
            bad.append((rel(p), str(e).strip()[:300] or type(e).__name__))
            continue
        for s, _, _ in g:
            origin[s].add(rel(p))
        data += g
    for f, err in bad:
        print(f"validate: invalid Turtle in {f}: {err}\n  (hint: don't use '/' in prefixed names; write the full IRI <http://kb.local/rfi/id/...>)")
    if not len(data):
        if not quiet:
            print("no triples to validate")
        return not bad
    ont = rdflib.Graph()
    for op in ontology_files():
        ont.parse(op, format="turtle")
    shapes = rdflib.Graph()
    for sp in shapes_files():
        shapes.parse(sp, format="turtle")
    conforms, rg, _ = validate(data, shacl_graph=shapes, ont_graph=ont, inference="rdfs", abort_on_first=False, allow_warnings=True)
    conforms = conforms and not bad
    if quiet:
        return conforms
    SH = rdflib.Namespace(PREFIXES["sh"])
    res = []
    for r in rg.subjects(rdflib.RDF.type, SH.ValidationResult):
        focus = rg.value(r, SH.focusNode)
        res.append([short(str(focus)), short(str(rg.value(r, SH.resultPath) or "")), str(rg.value(r, SH.resultMessage) or ""),
                    ", ".join(sorted(origin.get(focus, [])))])
    if conforms:
        print(f"validate: ok ({len(data)} triples, no violations)")
    else:
        print(f"validate: {len(res)} violations")
        print(fmt_table(["node", "property", "message", "file"], sorted(res)[:200], 70))
    return conforms


def cmd_validate(args, cfg):
    ok = do_validate(cfg)
    sys.exit(0 if ok else 3)


def result_alerts(values: set[str], graphs: dict) -> list[str]:
    alerts = []
    for gname, meta in graphs.items():
        if meta.get("layer") != "wiki":
            continue
        if gname in values or meta["doc"] in values:
            stale, days = note_is_stale(meta)
            if stale:
                alerts.append(f"overdue note: {meta['source']} ({days} days overdue) — confirm with the author")
            if meta.get("header_inferred"):
                alerts.append(f"inferred header: {meta['source']} (date/author not declared)")
    return alerts


def cmd_query(args, cfg):
    q = args.sparql
    p = ROOT / q if not q.lstrip().upper().startswith(("SELECT", "PREFIX", "ASK", "CONSTRUCT", "DESCRIBE")) else None
    if p is not None:
        if not p.exists():
            p = find_query(q)
        if p is None or not p.exists():
            die(f"query not found: {q} (workspace kb/queries/ or core {CORE_QUERIES})")
        q = p.read_text(encoding="utf-8")
    params = dict(kv.split("=", 1) for kv in (args.param or []))
    missing = sorted(set(re.findall(r"\{\{(\w+)\}\}", q)) - set(params))
    if missing:
        die("required parameters: " + ", ".join(f"--param {m}=..." for m in missing))
    for k, v in params.items():
        q = q.replace("{{" + k + "}}", v.replace("\\", "\\\\").replace('"', '\\"'))
    store, graphs = build_store()
    res = run_query(store, q)
    if isinstance(res, bool):
        print(res)
        return
    if not hasattr(res, "variables"):
        for t in res:
            print(t)
        return
    vars_ = [v.value for v in res.variables]
    rows, values = [], set()
    for sol in res:
        row = []
        for v in vars_:
            t = sol[v]
            s = term_str(t)
            values.add(s)
            row.append(s if args.full_iri else short(s))
        rows.append(row)
    if args.json:
        print(json.dumps([dict(zip(vars_, r)) for r in rows], ensure_ascii=False, indent=1))
    else:
        print(fmt_table(vars_, rows, args.width))
        print(f"\n{len(rows)} row(s)")
    for a in result_alerts(values, graphs):
        print(f"ALERT: {a}")


def find_subjects(store, term: str) -> list[str]:
    if term.startswith(("http://", "urn:")):
        return [term]
    if ":" in term.split()[0] and term.split(":")[0] in PREFIXES:
        return [expand(term)]
    esc = term.replace("\\", "\\\\").replace('"', '\\"')
    q = f"""SELECT DISTINCT ?s WHERE {{
      {{ ?s kb:originalId "{esc}" }} UNION {{ ?s kb:number "{esc}" }} UNION
      {{ ?s rdfs:label ?l FILTER(LCASE(STR(?l)) = LCASE("{esc}")) }} UNION {{ ?s skos:prefLabel ?l2 FILTER(LCASE(STR(?l2)) = LCASE("{esc}")) }}
    }} LIMIT 50"""
    return [term_str(s["s"]) for s in run_query(store, q)]


def entity_claims(store, s: str):
    q = f"SELECT ?p ?o ?g WHERE {{ GRAPH ?g {{ <{s}> ?p ?o }} }}"
    claims = defaultdict(lambda: defaultdict(set))
    for sol in run_query(store, q):
        claims[term_str(sol["p"])][term_str(sol["g"])].add(term_str(sol["o"]))
    return claims


def loc_of(store, s: str, g: str) -> str:
    q = f"""SELECT ?pg ?ln ?aba ?id WHERE {{ GRAPH <{g}> {{ OPTIONAL {{ <{s}> kb:page ?pg }} OPTIONAL {{ <{s}> kb:row ?ln }}
           OPTIONAL {{ <{s}> kb:sheet ?aba }} OPTIONAL {{ <{s}> kb:originalId ?id }} }} }} LIMIT 1"""
    for sol in run_query(store, q):
        parts = []
        if sol["pg"]:
            parts.append(f"p.{term_str(sol['pg'])}")
        if sol["aba"]:
            parts.append(f"sheet {term_str(sol['aba'])}")
        if sol["ln"]:
            parts.append(f"row {term_str(sol['ln'])}")
        return ", ".join(parts)
    return ""


def cmd_show(args, cfg):
    store, graphs = build_store()
    subs = find_subjects(store, args.term)
    if not subs:
        die(f"nothing found for '{args.term}' (try kb.py search)")
    for s in subs[: args.limit]:
        claims = entity_claims(store, s)
        print(f"\n## {short(s)}")
        used = set()
        rows = []
        for p, by_g in sorted(claims.items()):
            gs = {g: v for g, v in by_g.items() if g in graphs}
            if not gs:
                continue
            divergent = len({frozenset(v) for v in gs.values()}) > 1 and p not in PROV_PROPS
            winners = resolve(gs, graphs, s) if divergent else set(gs)
            for g, vals in gs.items():
                used.add(g)
                for v in sorted(vals):
                    mark = ("✓" if g in winners else "✗") if divergent else ""
                    rows.append([short(p), short(v), mark, f"{graphs[g]['source']} {loc_of(store, s, g)}".strip(), graphs[g].get("layer", "")])
        print(fmt_table(["property", "value", "wins", "source", "layer"], rows, args.width))
        for a in result_alerts(used | {graphs[g]["doc"] for g in used}, {g: graphs[g] for g in used}):
            print(f"ALERT: {a}")


STOPWORDS = set("""a o as os um uma uns umas de da do das dos em na no nas nos por para com sem sobre entre e ou que qual quais
quem como quando onde porque se ao aos à às é são ser foi está estão tem têm há isso este esta esse essa deve devem pelo pela
pelos pelas me nos lhe seu sua seus suas kb the of and or to in on for with what which who how when where is are
does do was were be by from that this these those el la los las del que cual quien como""".split())
TEXT_PROPS = {expand(p): w for p, w in (("rdfs:label", 3), ("skos:prefLabel", 3), ("kb:originalId", 3), ("kb:number", 3),
                                         ("dcterms:title", 2), ("kb:excerpt", 1),
                                         ("skos:definition", 1), ("kb:description", 1))}
ASK_SKIP_PROPS = WIKI_HIDDEN | {expand(p) for p in ("kb:inSection", "kb:parentSection", "rdfs:label", "skos:prefLabel")}


def graphs_subjects_ok(score: dict, s: str) -> set:
    """Only boost subjects that already matched at least one term (keeps class words from flooding the ranking)."""
    return {s} if score.get(s) else set()


def best_graph(gs: list[str], graphs: dict, s: str) -> str:
    """Source to cite: precedence winner; tie (identical copies) -> latest date/version/version name."""
    win = resolve({g: set() for g in gs}, graphs, s) or set(gs)
    return max(win, key=lambda g: (str(graphs[g].get("date") or ""), graphs[g].get("version_key") or "",
                                   str(graphs[g].get("version") or ""), graphs[g]["source"]))


def cmd_ask(args, cfg):
    """Search + ranking + precedence + citation in one call (compact output for Claude)."""
    store, graphs = build_store()
    q = fold(args.question)
    ids = set(re.findall(r"\b\d+(?:\.\d+)+\b", args.question))
    terms = [t for t in re.findall(r"[a-z0-9][a-z0-9./-]{2,}", q) if t not in STOPWORDS]
    if not terms and not ids:
        die("question has no search terms")
    text_props = dict(TEXT_PROPS)
    text_props.update({expand(k): int(v) for k, v in ((cfg.get("ask") or {}).get("text_props") or {}).items()})
    fontes = [fold(f) for f in (args.fonte or [])]
    if fontes:
        graphs = {g: m for g, m in graphs.items() if any(f in fold(m["source"]) for f in fontes)}
        if not graphs:
            die("no source matches --fonte " + ", ".join(args.fonte))
    rows = run_query(store, """SELECT ?s ?p ?v ?g WHERE { GRAPH ?g { ?s ?p ?v } FILTER(isLiteral(?v))
                               FILTER(?g != <urn:kb:ontology> && !STRSTARTS(STR(?s), "urn:kb:")) }""")
    score: dict[str, dict[str, int]] = defaultdict(dict)
    for r in rows:
        p = term_str(r["p"])
        w = text_props.get(p)
        if not w or term_str(r["g"]) not in graphs:
            continue
        s, v = term_str(r["s"]), fold(term_str(r["v"]))
        hits = score[s]
        for t in terms:
            if t in v:
                hits[t] = max(hits.get(t, 0), w)
        if p in (expand("kb:originalId"), expand("kb:number")) and term_str(r["v"]) in ids:
            hits["#id"] = 20
    # class names count as text too ("deadline" finds kb:Deadline/kb:Milestone instances, "requirement" finds requirements)
    cls_words = {"Milestone": "milestone deadline date marco prazo hito", "Deadline": "deadline prazo plazo date",
                 "Fact": "fact value", "RACIAssignment": "raci responsible accountable", "Term": "term glossary acronym"}
    for r in run_query(store, "SELECT DISTINCT ?s ?t WHERE { GRAPH ?g { ?s a ?t } FILTER(?g != <urn:kb:ontology>) }"):
        if term_str(r["s"]) not in graphs_subjects_ok(score, term_str(r["s"])):
            continue
        local = term_str(r["t"]).split("#")[-1]
        words = fold(re.sub(r"(?<!^)(?=[A-Z])", " ", local) + " " + cls_words.get(local, ""))
        for t in terms:
            if t in words.split():
                score[term_str(r["s"])][t] = max(score[term_str(r["s"])].get(t, 0), 2)
    exact = {term_str(r["s"]) for r in rows
             if term_str(r["p"]) in (expand("rdfs:label"), expand("skos:prefLabel")) and fold(term_str(r["v"])).strip(" ?.") == q.strip(" ?.")}

    def rank(s: str, h: dict) -> int:
        bonus = 5 if len(h) >= len(terms) else 0
        bonus += 4 if s in exact else 0
        bonus += 3 if s.startswith(NS_ID + "topic/") else 0  # topics consolidate sources and already carry precedence
        return sum(h.values()) + bonus

    ranked = sorted(((rank(s, h), s) for s, h in score.items() if h), reverse=True)
    if not ranked:
        print("NO ANSWER IN KB: no entity contains the terms" + (" in the filtered sources" if fontes else "")
              + ". If needed, delegate a full-text search to a subagent (Grep over the raw/ .md files + kb.py page-text).")
        return
    best_cov = max(len([t for t in score[s] if t != "#id"]) for _, s in ranked[: args.limit])
    if terms and best_cov < len(terms) and not any("#id" in score[s] for _, s in ranked[:3]):
        print(f"WARNING: no entity covers all terms ({best_cov}/{len(terms)}: {', '.join(terms)}). "
              "The answer may not be in the KB; do not conclude without confirming.\n")
    labels = {term_str(r["s"]): term_str(r["l"]) for r in run_query(store, "SELECT ?s (SAMPLE(?x) AS ?l) WHERE { { ?s rdfs:label ?x } UNION { ?s skos:prefLabel ?x } } GROUP BY ?s")}
    used = set()
    for i, (sc, s) in enumerate(ranked[: args.limit], 1):
        claims = entity_claims(store, s)
        gs = sorted({g for by_g in claims.values() for g in by_g if g in graphs})
        if not gs:
            continue
        g0 = best_graph(gs, graphs, s)
        used.update(gs)
        types = sorted({short(v).split(":", 1)[-1] for v in claims.get(expand("rdf:type"), {}).get(g0, set())})
        copies = len({graphs[g]["source"] for g in gs}) - 1
        print(f"[{i}] {labels.get(s, short(s))}  ({', '.join(types)})  {short(s)}")
        print(f"    source: {graphs[g0]['source']}" + (f", {loc_of(store, s, g0)}" if loc_of(store, s, g0) else "")
              + (f"  (+{copies} other source(s))" if copies else ""))
        for p, by_g in sorted(claims.items()):
            if p in ASK_SKIP_PROPS:
                continue
            gsp = {g: v for g, v in by_g.items() if g in graphs}
            if not gsp:
                continue
            divergent = len({frozenset(v) for v in gsp.values()}) > 1 and p not in PROV_PROPS
            win = resolve(gsp, graphs, s) if divergent else set(gsp)
            vals = sorted(set().union(*(gsp[g] for g in win)))
            lim = args.width * 3 if p == expand("kb:excerpt") else args.width
            shown = "; ".join(re.sub(r"\s+", " ", labels.get(v, short(v)) if v.startswith("http") else v)[:lim] for v in vals[:6])
            more = f" (+{len(vals) - 6})" if len(vals) > 6 else ""
            print(f"    {short(p).split(':', 1)[1]}: {shown}{more}")
            if divergent:
                for g, v in gsp.items():
                    if g not in win:
                        superseded = re.sub(r"\s+", " ", "; ".join(sorted(v)))[:160]
                        print(f"      ✗ superseded: {superseded} ({graphs[g]['source']})")
        if "Activity" in types:
            raci = list(run_query(store, f"""SELECT DISTINCT ?pl ?c WHERE {{ ?a kb:activity <{s}> ; kb:role ?p ; kb:raciCode ?c .
                                             ?p rdfs:label ?pl }} ORDER BY ?pl"""))
            if raci:
                print("    RACI: " + "; ".join(f"{term_str(r['pl'])}={term_str(r['c'])}" for r in raci))
        md = next((ROOT / graphs[g0]["source"]).with_name(Path(graphs[g0]["source"]).stem + ".md") for _ in [0])
        pg = claims.get(expand("kb:page"), {}).get(g0)
        if pg and md.exists():
            print(f'    literal text: {KB_CMD} page-text "{rel(md)}" {sorted(pg)[0]}')
    if len(ranked) > args.limit:
        print(f"\n(+{len(ranked) - args.limit} lower-scoring entity(ies); use --limit or more specific terms)")
    for a in result_alerts(used | {graphs[g]["doc"] for g in used}, {g: graphs[g] for g in used}):
        print(f"ALERT: {a}")


def cmd_search(args, cfg):
    store, graphs = build_store()
    esc = args.text.replace("\\", "\\\\").replace('"', '\\"')
    cls = f"?s a {args.cls} ." if args.cls else ""
    q = f"""SELECT DISTINCT ?s ?tipo ?p ?v ?g WHERE {{ GRAPH ?g {{ ?s ?p ?v . {cls} }}
            FILTER(?p IN (rdfs:label, kb:excerpt, skos:prefLabel, skos:definition, kb:originalId, dcterms:title))
            FILTER(CONTAINS(LCASE(STR(?v)), LCASE("{esc}")))
            OPTIONAL {{ ?s a ?tipo FILTER(STRSTARTS(STR(?tipo), "{NS_KB}") || STRSTARTS(STR(?tipo), "{NS_DOMAIN}")) }} }} LIMIT {args.limit}"""
    rows = []
    for sol in run_query(store, q):
        g = term_str(sol["g"])
        s = term_str(sol["s"])
        rows.append([short(s), short(term_str(sol["tipo"])), short(term_str(sol["p"])), term_str(sol["v"]),
                     f"{graphs.get(g, {}).get('source', g)} {loc_of(store, s, g)}".strip()])
    print(fmt_table(["entity", "type", "prop", "value", "source"], rows, args.width))
    print(f"\n{len(rows)} result(s)")


def cmd_stale(args, cfg):
    sources, _, _ = discover(cfg)
    rows = []
    for s in sources:
        if s.layer != "wiki":
            continue
        meta = doc_meta(s, cfg)
        stale, days = note_is_stale(meta)
        if stale or args.all:
            due = to_date(meta["date"]) + parse_duration(meta["review_every"], 14)
            rows.append((days, [s.rel, str(meta["date"]), meta["author"], str(meta["review_every"]), due.isoformat(),
                                str(days) if stale else "-", "yes" if meta["header_inferred"] else ""]))
    rows.sort(key=lambda x: -x[0])
    if not rows:
        print("no overdue notes")
        return
    print(fmt_table(["note", "last_reviewed", "reviewed_by", "review_every", "due", "days overdue", "inferred header"], [r for _, r in rows]))
    print("\n(alert only: precedence is unchanged; update last_reviewed/reviewed_by when reviewing)")


NOTE_TEMPLATES = {
    "en": ("Context", "Notes", "Decisions and hypotheses",
           "Graph conventions (one per line): - Decision: <text> | - Hypothesis: <text> | - Assumption: <text> | "
           "- Deadline: <milestone> = YYYY-MM-DD | - Term: <ACRONYM> = <definition> | - Condition: <topic> = <value>. "
           "Link notes with [[note-title]] and sources with [text](../raw/...). To override a raw/ document "
           "or topic explicitly, list it in `overrides:`."),
    "pt": ("Contexto", "Notas", "Decisões e hipóteses",
           "Convenções para o grafo (uma por linha): - Decisão: <texto> | - Hipótese: <texto> | - Premissa: <texto> | "
           "- Prazo: <marco> = AAAA-MM-DD | - Termo: <SIGLA> = <definição> | - Condição: <tema> = <valor>. "
           "Ligue notas com [[titulo-da-nota]] e fontes com [texto](../raw/...). Para sobrepor um documento de raw/ "
           "ou um tema explicitamente, liste-o em `overrides:`."),
    "es": ("Contexto", "Notas", "Decisiones e hipótesis",
           "Convenciones para el grafo (una por línea): - Decisión: <texto> | - Hipótesis: <texto> | - Supuesto: <texto> | "
           "- Plazo: <hito> = AAAA-MM-DD | - Término: <SIGLA> = <definición> | - Condición: <tema> = <valor>. "
           "Enlace notas con [[titulo-de-la-nota]] y fuentes con [texto](../raw/...). Para sobrescribir un documento "
           "de raw/ o un tema explícitamente, inclúyalo en `overrides:`."),
}


def cmd_new_note(args, cfg):
    wiki_dir = ROOT / cfg["sources"]["wiki"][0]
    p = wiki_dir / f"{today().isoformat()}-{slug(args.title)}.md"
    if p.exists():
        die(f"already exists: {rel(p)}")
    fm = {"title": args.title, "type": args.type, "last_reviewed": today(), "reviewed_by": current_user_name(),
          "review_every": cfg.get("review_every", "14d"), "overrides": []}
    if args.source:
        fm["source"] = args.source
    h_ctx, h_notes, h_dec, conv = NOTE_TEMPLATES.get((cfg.get("language") or "en")[:2], NOTE_TEMPLATES["en"])
    body = f"""
# {args.title}

## {h_ctx}

## {h_notes}

## {h_dec}
<!-- {conv} -->
"""
    wiki_dir.mkdir(parents=True, exist_ok=True)
    p.write_text(dump_frontmatter(fm) + body, encoding="utf-8")
    print(f"created: {rel(p)}")


# ----------------------------------------------------------------------------- browsable wiki


def do_wiki(cfg) -> int:
    store, graphs = build_store()
    files: dict[str, str] = {}
    label_q = """SELECT ?s (SAMPLE(?l) AS ?label) WHERE { { ?s rdfs:label ?l } UNION { ?s skos:prefLabel ?l } } GROUP BY ?s"""
    labels = {term_str(r["s"]): term_str(r["label"]) for r in run_query(store, label_q)}
    type_q = f"""SELECT ?s ?t WHERE {{ GRAPH ?g {{ ?s a ?t }} FILTER((STRSTARTS(STR(?t), "{NS_KB}") || STRSTARTS(STR(?t), "{NS_DOMAIN}")) && ?g != <urn:kb:ontology>) }}"""
    types = defaultdict(set)
    for r in run_query(store, type_q):
        t = term_str(r["t"])
        types[term_str(r["s"])].add(t.split("#", 1)[-1])
    inline = {"Section", "RACIAssignment", "Document", "WikiNote"}

    def page_path(s: str) -> str | None:
        if s.startswith(NS_ID + "doc/") or s.startswith(NS_ID + "note/"):
            return "documents/" + s.rsplit("/", 1)[1] + ".md"
        if not s.startswith(NS_ID) or (types.get(s, set()) & inline) or not types.get(s):
            return None
        return "e/" + s[len(NS_ID):] + ".md"

    def link(s: str, frm: str) -> str:
        lbl = (labels.get(s) or short(s)).replace("|", "\\|").replace("\n", " ")
        pp = page_path(s)
        if not pp:
            return lbl
        relp = os.path.relpath(pp, os.path.dirname(frm) or ".")
        return f"[{lbl[:90]}]({relp.replace(' ', '%20')})"

    def val(v: str, frm: str) -> str:
        if v.startswith(NS_ID):
            return link(v, frm)
        if v.startswith("http"):
            return short(v)
        return v.replace("|", "\\|").replace("\n", " ")[:300]

    # entity pages
    by_class = defaultdict(list)
    incoming = defaultdict(set)
    for r in run_query(store, f"""SELECT DISTINCT ?s ?o WHERE {{ ?s ?p ?o FILTER(isIRI(?o) && STRSTARTS(STR(?o), "{NS_ID}") && ?p != kb:definedIn) }}"""):
        incoming[term_str(r["o"])].add(term_str(r["s"]))
    for s, ts in types.items():
        pp = page_path(s)
        if not pp or not pp.startswith("e/"):
            continue
        for t in ts:
            by_class[t].append(s)
        claims = entity_claims(store, s)
        lines = [f"# {labels.get(s, short(s))}", "", f"- IRI: `{short(s)}`", f"- Type: {', '.join(sorted(ts))}", "",
                 "| property | value | source | layer | wins |", "|---|---|---|---|---|"]
        used = set()
        for p, by_g in sorted(claims.items()):
            gs = {g: v for g, v in by_g.items() if g in graphs}
            divergent = len({frozenset(v) for v in gs.values()}) > 1 and p not in PROV_PROPS
            winners = resolve(gs, graphs, s) if divergent else set(gs)
            for g, vals in sorted(gs.items()):
                used.add(g)
                for v in sorted(vals):
                    if p in WIKI_HIDDEN:  # already shown in the "source" column/header
                        continue
                    lines.append(f"| {short(p)} | {val(v, pp)} | {link(graphs[g]['doc'], pp)} {loc_of(store, s, g)} | {graphs[g].get('layer', '')} | "
                                 f"{('✓' if g in winners else '✗') if divergent else ''} |")
        for a in result_alerts(used | {graphs[g]["doc"] for g in used}, {g: graphs[g] for g in used}):
            lines.append(f"\n> ⚠️ {a}")
        refs = sorted(incoming.get(s, set()))
        if refs:
            lines += ["", "## Referenced by", ""] + [f"- {link(x, pp)}" for x in refs[:100]]
            # inline RACI assignments
        raci = [x for x in refs if "RACIAssignment" in types.get(x, set())]
        if raci:
            lines += ["", "## RACI", "", "| role | code |", "|---|---|"]
            for a in raci:
                rc = entity_claims(store, a)
                papel = next(iter(next(iter(rc.get(expand("kb:role"), {}).values()), set())), "")
                code = next(iter(next(iter(rc.get(expand("kb:raciCode"), {}).values()), set())), "")
                lines.append(f"| {link(papel, pp)} | {code} |")
        files[pp] = "\n".join(lines) + "\n"

    # document/note pages
    docs = []
    for gname, meta in graphs.items():
        if meta["method"] != "det":
            continue
        d = meta["doc"]
        pp = page_path(d)
        docs.append((meta, d, pp))
        secs = list(run_query(store, f"""SELECT ?s ?l ?pg ?num WHERE {{ GRAPH <{gname}> {{ ?s a kb:Section ; rdfs:label ?l .
                     OPTIONAL {{ ?s kb:page ?pg }} OPTIONAL {{ ?s kb:number ?num }} }} }}"""))
        secs.sort(key=lambda r: (int(term_str(r["pg"]) or 0), term_str(r["l"])))
        counts = list(run_query(store, f"""SELECT ?t (COUNT(DISTINCT ?s) AS ?n) WHERE {{ GRAPH <{gname}> {{ ?s a ?t ; kb:definedIn <{d}> }}
                       FILTER(?t != kb:Section) }} GROUP BY ?t ORDER BY DESC(?n)"""))
        stale, days = note_is_stale(meta)
        lines = [f"# {meta.get('title') or meta['source']}", "",
                 f"- Source: `{meta['source']}`", f"- Layer: {meta.get('layer')}", f"- Date: {meta.get('date') or '?'}",
                 f"- Version: {meta.get('version') or '-'}", f"- Family: {meta.get('family') or '-'}"]
        if meta.get("layer") == "wiki":
            lines += [f"- Reviewed by: {meta.get('author')}", f"- Review every: {meta.get('review_every')}"]
            if meta.get("overrides"):
                lines.append("- Overrides: " + ", ".join(link(o, pp) for o in meta["overrides"]))
            if stale:
                lines.append(f"\n> ⚠️ review overdue by {days} days")
            if meta.get("header_inferred"):
                lines.append("\n> ⚠️ inferred header (date/author not declared)")
        for w in meta.get("warnings") or []:
            lines.append(f"\n> warning: {w}")
        if counts:
            lines += ["", "## Extracted entities", ""] + [f"- {short(term_str(r['t']))}: {term_str(r['n'])}" for r in counts]
        if secs:
            lines += ["", "## Sections", ""]
            for r in secs:
                pg = term_str(r["pg"])
                lines.append(f"- {term_str(r['l'])}" + (f" (p.{pg})" if pg else ""))
        files[pp] = "\n".join(lines) + "\n"

    # index by class
    for cls, subs in by_class.items():
        pp = f"classes/{cls}.md"
        subs = sorted(set(subs), key=lambda x: (labels.get(x) or x).lower())
        files[pp] = f"# {cls} ({len(subs)})\n\n" + "\n".join(f"- {link(x, pp)}" for x in subs) + "\n"

    # conflicts
    conf_lines = ["# Conflicts between sources", "", "Divergent values for the same subject/property. ✓ = precedence winner",
                  "(most recent wins; raw beats wiki, unless `overrides`).", "", "| entity | property | value | source | wins |", "|---|---|---|---|---|"]
    qp = find_query("conflitos") or find_query("conflicts")
    q = qp.read_text(encoding="utf-8") if qp else None
    n_conf = 0
    if q:
        pairs = {(term_str(r["s"]), term_str(r["p"])) for r in run_query(store, q)}
        for s, p in sorted(pairs):
            by_g = entity_claims(store, s).get(p, {})
            gs = {g: v for g, v in by_g.items() if g in graphs}
            winners = resolve(gs, graphs, s)
            n_conf += 1
            for g, vals in sorted(gs.items()):
                for v in sorted(vals):
                    conf_lines.append(f"| {link(s, 'conflicts.md')} | {short(p)} | {val(v, 'conflicts.md')} | {graphs[g]['source']} | {'✓' if g in winners else '✗'} |")
    if not n_conf:
        conf_lines.append("| _(none)_ | | | | |")
    files["conflicts.md"] = "\n".join(conf_lines) + "\n"

    # pending reviews
    rv = ["# Pending reviews (wiki/)", "", "| note | last reviewed | by | review every | overdue (days) | inferred header |", "|---|---|---|---|---|---|"]
    for meta, d, pp in sorted(docs, key=lambda x: str(x[0].get("date"))):
        if meta.get("layer") != "wiki":
            continue
        stale, days = note_is_stale(meta)
        if stale or meta.get("header_inferred"):
            rv.append(f"| {link(d, 'pending-reviews.md')} | {meta.get('date')} | {meta.get('author')} | {meta.get('review_every')} | {days if stale else '-'} | {'yes' if meta.get('header_inferred') else ''} |")
    files["pending-reviews.md"] = "\n".join(rv) + "\n"

    # index
    idx = [f"# KB — {cfg.get('title') or ROOT.name}", "", "Generated by `kb.py wiki` from the graph. Do not edit by hand.", "",
           "- [Conflicts between sources](conflicts.md)", "- [Pending reviews](pending-reviews.md)", "",
           "## Documents (raw/)", "", "| document | version | date | family |", "|---|---|---|---|"]
    for meta, d, pp in sorted(docs, key=lambda x: x[0]["source"]):
        if meta.get("layer") == "raw":
            idx.append(f"| {link(d, 'index.md')} | {meta.get('version') or '-'} | {meta.get('date') or '?'} | {meta.get('family')} |")
    idx += ["", "## Notes (wiki/)", "", "| note | last reviewed | by |", "|---|---|---|"]
    for meta, d, pp in sorted(docs, key=lambda x: x[0]["source"]):
        if meta.get("layer") == "wiki":
            idx.append(f"| {link(d, 'index.md')} | {meta.get('date')} | {meta.get('author')} |")
    idx += ["", "## Classes", ""] + [f"- [{c}](classes/{c}.md) ({len(set(v))})" for c, v in sorted(by_class.items())]
    files["index.md"] = "\n".join(idx) + "\n"

    changed = 0
    for pth, content in files.items():
        if write_if_changed(WIKI_OUT / pth, content):
            changed += 1
    if WIKI_OUT.exists():
        for f in WIKI_OUT.rglob("*.md"):
            if f.relative_to(WIKI_OUT).as_posix() not in files:
                f.unlink()
                changed += 1
        for d in sorted((p for p in WIKI_OUT.rglob("*") if p.is_dir()), key=lambda p: -len(p.parts)):
            if not any(d.iterdir()):
                d.rmdir()
    return changed


def cmd_wiki(args, cfg):
    print(f"wiki: {do_wiki(cfg)} files changed in kb/wiki/")



# ----------------------------------------------------------------------------- lifecycle: scaffold / migrate / volume

WRITE_CMDS = {"log-cost", "convert", "extract-det", "update", "ingest", "lock", "unlock", "unzip", "set-image", "llm-done", "wiki",
              "new-note", "migrate", "scaffold", "vendor"}
GITIGNORE_BLOCK = """
# --- kb-setup ---
.DS_Store
~$*
*.tmp
# binaries stay out of git; their .md conversions are versioned
*.pdf
*.docx
*.xlsx
*.xlsm
*.zip
raw/**/*.assets/
kb/.lock
graphify-out/
# --- /kb-setup ---
"""


def cmd_scaffold(args, cfg):
    """Create the Zettelkasten workspace skeleton (raw/ wiki/ kb/ docs/) and kb/config.yaml in the current folder."""
    if getattr(args, "root", None):
        root = Path(args.root).expanduser().resolve()
    elif os.environ.get("KB_ROOT"):
        root = Path(os.environ["KB_ROOT"]).expanduser().resolve()
    else:
        root = Path.cwd().resolve()  # scaffold never walks up: the target is the current folder
    cfg_path = root / "kb" / "config.yaml"
    if cfg_path.exists() and not args.force:
        die(f"{cfg_path} already exists (use adopt/upgrade, or --force)")
    if IS_COPY or not (TEMPLATES / "config.yaml").exists():
        die(f"scaffold needs the kb-setup templates, and this engine is a workspace copy ({ENGINE_DIR}). "
            f"Install the kb-setup skill ({INSTALL_CMD}) and run {CENTRAL_CMD} scaffold")
    tpl = (TEMPLATES / "config.yaml").read_text(encoding="utf-8")
    ws = re.sub(r"[^a-z0-9]+", "-", fold(args.title or root.name)).strip("-")[:40] or "ws"
    prefix = args.prefix or "dom"
    values = {"engine_version": ENGINE_VERSION, "kb_setup_version": KB_SETUP_VERSION, "title": args.title or root.name, "language": args.language,
              "prefix": prefix, "namespace": f"http://kb.local/{ws}/ont#", "instances": f"http://kb.local/{ws}/id/",
              "versioning": args.versioning}
    for k, v in values.items():
        tpl = tpl.replace("{{" + k + "}}", str(v))
    created = []
    for d in ("raw", "wiki", "docs", "kb/ontology", "kb/mappings", "kb/queries", "kb/evals"):
        if not (root / d).exists():
            (root / d).mkdir(parents=True)
            created.append(d + "/")
    cfg_path.write_text(tpl, encoding="utf-8")
    created.append("kb/config.yaml")
    imp = root / "kb" / "IMPROVEMENTS.md"
    if not imp.exists():
        imp.write_text((TEMPLATES / "IMPROVEMENTS.md").read_text(encoding="utf-8"), encoding="utf-8")
        created.append("kb/IMPROVEMENTS.md")
    if args.versioning in ("git", "both"):
        gi = root / ".gitignore"
        cur = gi.read_text(encoding="utf-8") if gi.exists() else ""
        if "# --- kb-setup ---" not in cur:
            gi.write_text(cur.rstrip("\n") + "\n" + GITIGNORE_BLOCK, encoding="utf-8")
            created.append(".gitignore (kb-setup block)")
    print("scaffold: " + ", ".join(created))


def _v1_namespaces() -> tuple[str, str] | None:
    for p in sorted(ONTOLOGY_DIR.glob("*.ttl")) if ONTOLOGY_DIR.exists() else []:
        for m in re.finditer(r"@prefix\s+(\w+):\s*<(http://kb\.local/[^>]+/ont#)>", p.read_text(encoding="utf-8")):
            if m.group(2) != NS_KB:
                return m.group(1), m.group(2)
    return None


def _rename_terms(text: str, prefix: str, ns: str, renames: dict, id_ns: str) -> str:
    """v1 -> v2: prefixed core terms (rfi:definidoEm -> kb:definedIn), full IRIs and instance path segments."""
    for pt, en in sorted(renames.items(), key=lambda kv: -len(kv[0])):
        text = re.sub(rf"(?<![\w:]){re.escape(prefix)}:{pt}\b", f"kb:{en}", text)
        text = text.replace(f"<{ns}{pt}>", f"<{NS_KB}{en}>")
    for a, b in V1_ID_SEGMENTS.items():
        text = text.replace(f"{id_ns}{a}/", f"{id_ns}{b}/").replace(f"id:{a}/", f"id:{b}/")
    return text


V1_ID_SEGMENTS = {"nota": "note", "secao": "section", "termo": "term", "marco": "milestone", "prazo": "deadline",
                  "condicao": "topic", "pessoa": "person", "papel": "role", "decisao": "decision", "hipotese": "hypothesis",
                  "premissa": "assumption"}


def _ensure_prefix_decl(text: str, kind: str) -> str:
    if kind == "ttl" and "@prefix kb:" not in text:
        return f"@prefix kb:      <{NS_KB}> .\n" + text
    if kind == "rq" and "PREFIX kb:" not in text:
        m = re.search(r"^PREFIX\s", text, re.M)
        line = f"PREFIX kb: <{NS_KB}>\n"
        return text[: m.start()] + line + text[m.start():] if m else line + text
    return text


def _migrate_config_text(text: str, prefix: str, ns: str, id_ns: str, title: str, language: str) -> str:
    sys.path.insert(0, str(ENGINE_DIR))
    import core_renames as R

    out, section = [], None
    for line in text.splitlines():
        top = re.match(r"^(\w+):", line)
        if top:
            section = R.CONFIG_KEYS.get(top.group(1), top.group(1))
            line = re.sub(r"^\w+:", section + ":", line, count=1)
        else:
            keymap = {"topics": R.TOPIC_KEYS, "facts": R.FACT_KEYS, "attachments": R.ATTACHMENT_KEYS}.get(section, {})
            m = re.match(r"^(\s+-?\s*)(\w+):", line)
            if m and m.group(2) in keymap:
                line = m.group(1) + keymap[m.group(2)] + ":" + line[m.end():]
            if section == "facts":
                mm = re.match(r"^(\s+normalize:\s*)(\w+)(.*)$", line)
                if mm and mm.group(2) in R.FACT_NORMALIZE:
                    line = mm.group(1) + R.FACT_NORMALIZE[mm.group(2)] + mm.group(3)
        out.append(line)
    body = "\n".join(out) + "\n"
    body = _rename_terms(body, prefix, ns, R.ALL, id_ns)
    head = (f"# Migrated to kb-setup engine {ENGINE_VERSION} on {today().isoformat()} (v1 -> v2).\n"
            f"engine_version: \"{ENGINE_VERSION}\"\nkb_setup_version: \"{KB_SETUP_VERSION}\"\ntitle: \"{title}\"\nlanguage: {language}\n"
            f"ontology:\n  prefix: {prefix}\n  namespace: \"{ns}\"\n  instances: \"{id_ns}\"\n\n")
    return head + body


def cmd_migrate(args, cfg):
    """Adopt a v1 workspace (Portuguese core terms, local scripts/kb/kb.py) into the v2 central engine."""
    sys.path.insert(0, str(ENGINE_DIR))
    import core_renames as R

    if (cfg.get("engine_version") or "").startswith("2") and not args.force:
        die(f"workspace already on engine {cfg.get('engine_version')}")
    found = _v1_namespaces()
    if not found and not args.prefix:
        die("could not detect the v1 domain namespace in kb/ontology/*.ttl; pass --prefix and --namespace")
    prefix, ns = (args.prefix or found[0]), (args.namespace or found[1])
    id_ns = ns.replace("/ont#", "/id/")
    plan: list[tuple[Path, str]] = []

    def stage(p: Path, new: str):
        if p.read_text(encoding="utf-8") != new:
            plan.append((p, new))

    cp = KB / "config.yaml"
    stage(cp, _migrate_config_text(cp.read_text(encoding="utf-8"), prefix, ns, id_ns, args.title or ROOT.name, args.language))
    for p in sorted(MAPPINGS_DIR.glob("*.yaml")) if MAPPINGS_DIR.exists() else []:
        stage(p, _rename_terms(p.read_text(encoding="utf-8"), prefix, ns, R.ALL, id_ns))
    for p in sorted(QUERIES_DIR.glob("*.rq")) if QUERIES_DIR.exists() else []:
        stage(p, _ensure_prefix_decl(_rename_terms(p.read_text(encoding="utf-8"), prefix, ns, R.ALL, id_ns), "rq"))
    for p in sorted(ONTOLOGY_DIR.glob("*.ttl")) if ONTOLOGY_DIR.exists() else []:
        stage(p, _ensure_prefix_decl(_rename_terms(p.read_text(encoding="utf-8"), prefix, ns, R.ALL, id_ns), "ttl"))
    llm_dir = TRIPLES_DIR / "llm"
    for p in sorted(llm_dir.glob("*.ttl")) if llm_dir.exists() else []:
        stage(p, _ensure_prefix_decl(_rename_terms(p.read_text(encoding="utf-8"), prefix, ns, R.ALL, id_ns), "ttl"))
    wiki_dir = ROOT / (cfg["sources"]["wiki"][0] if cfg.get("sources") else "wiki")
    for p in sorted(wiki_dir.rglob("*.md")) if wiki_dir.exists() else []:
        t = p.read_text(encoding="utf-8")
        nt = t
        for a, b in V1_ID_SEGMENTS.items():
            nt = nt.replace(f"id:{a}/", f"id:{b}/")
        if nt != t:
            plan.append((p, nt))
    print(f"migrate v1 -> v2 (prefix {prefix}: {ns}; instances {id_ns}); {len(plan)} file(s) to rewrite:")
    for p, _ in plan:
        print(f"  - {rel(p)}")
    det = sorted((TRIPLES_DIR / "det").glob("*.ttl")) if (TRIPLES_DIR / "det").exists() else []
    print(f"  - kb/triples/det/*.ttl: {len(det)} file(s) deleted and regenerated on the next update")
    if args.dry_run:
        print("dry-run: nothing written")
        return
    with CuratorLock(cfg, "migrate"):
        for p, new in plan:
            p.write_text(new, encoding="utf-8")
        for p in det:
            p.unlink()
    print("done. Next: review kb/ontology (drop declarations now provided by the core: "
          f"{CORE_ONTOLOGY}), run `update`, regenerate the /kb skill and CLAUDE.md, then remove the local scripts/kb/ engine.")


def cmd_volume(args, cfg):
    """Corpus volume (files, pages, words, MB) per source and totals; used by evals to relate cost to data size."""
    sources, _, _ = discover(cfg)
    sel = select(sources, args.path) if args.path else sources
    rows, tot = [], Counter()
    for s in sel:
        b = s.path.stat().st_size if s.binary_present else 0
        text = s.md.read_text(encoding="utf-8") if s.md.exists() else ""
        _, body = split_frontmatter(text)
        pages = len(re.findall(r"<!-- page: \d+ -->", body)) or None
        words = len(re.findall(r"\w+", body))
        imgs = len(parse_image_blocks(body)) if s.kind not in ("note", "md") else 0
        rows.append({"source": s.rel, "layer": s.layer, "kind": s.kind, "bytes": b, "md_chars": len(body),
                     "words": words, "pages": pages, "images": imgs, "est_tokens": int(len(body) / 4)})
        tot.update({"files": 1, "bytes": b, "md_chars": len(body), "words": words, "pages": pages or 0,
                    "images": imgs, "est_tokens": int(len(body) / 4)})
    out = {"root": ROOT.name, "sources": rows, "totals": dict(tot)}
    if args.json:
        print(json.dumps(out, ensure_ascii=False, indent=1))
        return
    print(fmt_table(["source", "kind", "MB", "pages", "words", "est_tokens"],
                    [[r["source"], r["kind"], f"{r['bytes'] / 1e6:.2f}", str(r["pages"] or "-"), str(r["words"]),
                      str(r["est_tokens"])] for r in rows], 70))
    print(f"\ntotal: {tot['files']} files, {tot['bytes'] / 1e6:.1f} MB, {tot['pages']} pages, {tot['words']} words, "
          f"~{tot['est_tokens']} tokens of converted text")


COSTS_FILE = KB / "evals" / "build-costs.jsonl"


def cmd_log_cost(args, cfg):
    """Record the token/time cost of an LLM build activity (image descriptions, extract-llm, graphify build, eval keys)."""
    COSTS_FILE.parent.mkdir(parents=True, exist_ok=True)
    rec = {"date": now_iso(), "kind": args.kind, "config": args.config, "tokens": args.tokens,
           "duration_ms": args.duration_ms, "tool_uses": args.tool_uses, "note": args.note or ""}
    with open(COSTS_FILE, "a", encoding="utf-8") as f:
        f.write(json.dumps(rec, ensure_ascii=False) + "\n")
    print(f"cost logged: {args.kind} {args.tokens} tokens")


def cmd_costs(args, cfg):
    rows = [json.loads(l) for l in COSTS_FILE.read_text(encoding="utf-8").splitlines() if l.strip()] if COSTS_FILE.exists() else []
    if args.json:
        print("\n".join(json.dumps(r, ensure_ascii=False) for r in rows))
        return
    tot = defaultdict(lambda: [0, 0])
    for r in rows:
        tot[(r["kind"], r["config"])][0] += r.get("tokens", 0)
        tot[(r["kind"], r["config"])][1] += r.get("duration_ms", 0)
    print(fmt_table(["kind", "config", "tokens", "minutes"],
                    [[k, c, f"{t:,}", f"{ms_ / 60000:.1f}"] for (k, c), (t, ms_) in sorted(tot.items())]) if tot else "no costs logged")


# ----------------------------------------------------------------------------- engine copy / versions

VENDOR_DIR = ROOT / "scripts" / "kb-engine"   # not scripts/kb/: that path marks a v1 workspace for adopt
UPDATE_URL = "https://raw.githubusercontent.com/wagnerpinheiro/graphify-kb/v8/kb-setup/VERSION"
CHANGELOG_URL = "https://github.com/wagnerpinheiro/graphify-kb/blob/v8/kb-setup/CHANGELOG.md"


def _semver(v: str) -> tuple[int, ...]:
    return tuple(int(x) for x in re.findall(r"\d+", str(v))[:3])


def _vendor_files() -> list[tuple[Path, str]]:
    """(source, path inside the copy) of every file a workspace copy of the engine needs to consult the KB."""
    files = [(ENGINE_DIR / n, n) for n in ("kb.py", "kb.py.lock", "core_renames.py")]
    files.append((_VERSION_FILE, "VERSION"))
    for sub, pat in (("assets/core", "*.ttl"), ("assets/queries", "*.rq")):
        files += [(p, f"{sub}/{p.name}") for p in sorted(_bundle(sub).glob(pat))]
    files.append((PROMPTS, "assets/templates/prompts.md"))
    return files


def vendor_state() -> tuple[str, str]:
    """State of scripts/kb-engine/: absent | up to date | outdated (vs the running engine) | modified (vs its ENGINE.json)."""
    meta_p = VENDOR_DIR / "ENGINE.json"
    if not meta_p.exists():
        return "absent", ""
    try:
        meta = json.loads(meta_p.read_text(encoding="utf-8"))
    except (OSError, ValueError):
        return "outdated", "unreadable ENGINE.json"
    have = meta.get("kb_setup_version") or "?"
    recorded = meta.get("files") or {}
    if IS_COPY:  # no central engine to compare with: check the copy against its own record
        diff = [r for r, h in recorded.items() if not (VENDOR_DIR / r).exists() or sha256_file(VENDOR_DIR / r) != h]
        if diff:
            return "modified", f"{len(diff)} file(s) differ from ENGINE.json: {', '.join(sorted(diff)[:5])}"
        return "up to date", f"kb-setup {have}, engine {meta.get('engine_version')}; integrity check only, run the central engine to compare"
    want = {r: sha256_file(src) for src, r in _vendor_files() if src.exists()}
    diff = sorted(r for r in want if not (VENDOR_DIR / r).exists() or sha256_file(VENDOR_DIR / r) != want[r])
    diff += sorted(r for r in recorded if r not in want)  # files the running engine no longer ships
    if not diff:
        return "up to date", f"kb-setup {have}, engine {meta.get('engine_version')}"
    return "outdated", f"kb-setup {have} → {KB_SETUP_VERSION}; {len(diff)} file(s) differ"


def cmd_vendor(args, cfg):
    """Copy the engine into scripts/kb-engine/ (text files only) so the workspace can be consulted without kb-setup."""
    if args.check:
        st, detail = vendor_state()
        print(f"scripts/kb-engine: {st}" + (f" ({detail})" if detail else ""))
        return
    if IS_COPY:
        die(f"this engine is already a workspace copy; run from the central engine: {CENTRAL_CMD} vendor")
    if not (KB / "config.yaml").exists():
        die(f"no kb/config.yaml in {ROOT}: run vendor from a workspace")
    if VENDOR_DIR.exists() and any(VENDOR_DIR.iterdir()) and not (VENDOR_DIR / "ENGINE.json").exists():
        die(f"{rel(VENDOR_DIR)}/ exists and is not an engine copy (no ENGINE.json): move it away first")
    files = _vendor_files()
    missing = [str(src) for src, _ in files if not src.exists()]
    if missing:
        die("engine files missing: " + ", ".join(missing))
    with CuratorLock(cfg, "vendor"):
        tmp = VENDOR_DIR.with_name(".kb-engine.tmp")
        if tmp.exists():
            shutil.rmtree(tmp)
        hashes = {}
        for src, r in files:
            dst = tmp / r
            dst.parent.mkdir(parents=True, exist_ok=True)
            shutil.copy2(src, dst)
            hashes[r] = sha256_file(dst)
        home = str(Path.home())
        origin = str(ENGINE_DIR)
        meta = {"engine_version": ENGINE_VERSION, "kb_setup_version": KB_SETUP_VERSION,
                "source": "~" + origin[len(home):] if origin.startswith(home) else origin,
                "vendored": now_iso(), "files": hashes}
        (tmp / "ENGINE.json").write_text(json.dumps(meta, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
        if VENDOR_DIR.exists():
            shutil.rmtree(VENDOR_DIR)
        tmp.rename(VENDOR_DIR)
    notes = []
    gi = ROOT / ".graphifyignore"
    if gi.exists() and "scripts/kb-engine" not in gi.read_text(encoding="utf-8"):
        gi.write_text(gi.read_text(encoding="utf-8").rstrip("\n") + "\nscripts/kb-engine/\n", encoding="utf-8")
        notes.append(".graphifyignore += scripts/kb-engine/")
    print(f"vendor: {rel(VENDOR_DIR)}/ ← engine {ENGINE_VERSION} · kb-setup {KB_SETUP_VERSION} ({len(files)} files + ENGINE.json)"
          + (f"; {'; '.join(notes)}" if notes else ""))
    print(f"consult without kb-setup: uv run {rel(VENDOR_DIR)}/kb.py status")


def cmd_check_update(args, cfg):
    """Compare the installed kb-setup VERSION with the published one (network, 3 s timeout; never fails)."""
    if IS_COPY:
        print("update check skipped: running from a workspace copy (the curator refreshes it with /kb-setup upgrade)")
        return
    import urllib.request

    url = os.environ.get("KB_SETUP_UPDATE_URL") or UPDATE_URL
    try:
        with urllib.request.urlopen(url, timeout=3) as r:
            latest = r.read(64).decode("utf-8", "replace").strip()
        if not re.fullmatch(r"\d+\.\d+\.\d+", latest):
            raise ValueError(f"unexpected content at {url}")
    except Exception as e:  # offline, proxy, 404, bad URL: never block the caller
        print(f"update check skipped: {e}")
        return
    if _semver(latest) <= _semver(KB_SETUP_VERSION):
        print(f"kb-setup {KB_SETUP_VERSION} is up to date (published: {latest})")
        return
    repo = SKILL_DIR.parent
    how = f"git -C {repo} pull" if (repo / ".git").exists() else "npx skills update kb-setup -g"
    print(f"UPDATE AVAILABLE: kb-setup {latest} (installed {KB_SETUP_VERSION})")
    print(f"  changelog: {CHANGELOG_URL}")
    print(f"  update:    {how}")
    print("  then open a new Claude Code session and run /kb-setup upgrade in each workspace")


def version_warnings(cfg: dict) -> list[str]:
    if not (KB / "config.yaml").exists():
        return []
    v = str(cfg.get("engine_version") or "")
    if not v:
        return [f"workspace has no engine_version (v1 layout?): run /kb-setup adopt (engine {ENGINE_VERSION})"]
    out = []
    if v.split(".")[:2] != ENGINE_VERSION.split(".")[:2]:
        out.append(f"workspace built with engine {v}, installed engine is {ENGINE_VERSION}: run /kb-setup upgrade")
    sv = str(cfg.get("kb_setup_version") or "")
    if sv and KB_SETUP_VERSION != "unknown" and _semver(sv)[:2] < _semver(KB_SETUP_VERSION)[:2]:
        out.append(f"workspace files generated by kb-setup {sv}; installed {KB_SETUP_VERSION}: run /kb-setup upgrade")
    return out

# ----------------------------------------------------------------------------- main


def main():
    ap = argparse.ArgumentParser(prog="kb.py", description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--root", help="workspace root (default: nearest ancestor with kb/config.yaml)")
    ap.add_argument("--read-only", action="store_true", help="consult-only mode: refuse commands that write (also KB_READONLY=1)")
    ap.add_argument("--version", action="version",
                    version=f"kb.py engine {ENGINE_VERSION} · kb-setup {KB_SETUP_VERSION} ({'workspace copy' if IS_COPY else 'central'})")
    sp = ap.add_subparsers(dest="cmd", required=True)

    def add(name, fn, help_):
        p = sp.add_parser(name, help=help_)
        p.set_defaults(fn=fn)
        return p

    add("status", cmd_status, "pending steps, lock, overdue notes, OneDrive conflict copies")
    p = add("unzip", cmd_unzip, "extract raw/ .zip files into a sibling folder")
    p.add_argument("--force", action="store_true")
    for name, fn, h in (("convert", cmd_convert, "convert raw/ binaries to .md (+ images in .assets/)"),
                        ("extract-det", cmd_extract_det, "generate deterministic triples in kb/triples/det/")):
        p = add(name, fn, h)
        p.add_argument("path", nargs="?")
        p.add_argument("--force", action="store_true")
        p.add_argument("-v", "--verbose", action="store_true")
    for name, fn, h in (("update", cmd_update, "unzip → convert → extract-det → validate → wiki, only what changed"),
                        ("ingest", cmd_ingest, "like update, optionally restricted to a path")):
        p = add(name, fn, h)
        p.add_argument("path", nargs="?")
        p.add_argument("--force", action="store_true")
    p = add("lock", cmd_lock, "acquire the curator lock (kb/.lock)")
    p.add_argument("--command-name", default=None)
    p = add("unlock", cmd_unlock, "release the curator lock")
    p.add_argument("--force", action="store_true")
    p = add("pending-images", cmd_pending_images, "list images with a PENDENTE marker")
    p.add_argument("path", nargs="?")
    p.add_argument("--json", action="store_true")
    p = add("image-context", cmd_image_context, "page/surrounding text of an image")
    p.add_argument("md")
    p.add_argument("asset")
    p = add("page-text", cmd_page_text, "text of one page of a PDF-converted .md")
    p.add_argument("md")
    p.add_argument("page", type=int)
    p = add("set-image", cmd_set_image, "write an image description (stdin or --text) and mark it DESCRITA")
    p.add_argument("md")
    p.add_argument("asset")
    p.add_argument("--text")
    p.add_argument("--ignore", nargs="?", const="", default=None, help="mark IGNORADA (decorative), with an optional reason")
    p = add("outline", cmd_outline, "headings/sections of the .md files and spreadsheet headers")
    p.add_argument("path", nargs="?")
    p = add("llm-context", cmd_llm_context, "paths and IRIs for the prose extraction of a source")
    p.add_argument("source")
    p = add("llm-done", cmd_llm_done, "record the llm extraction of a source in the manifest")
    p.add_argument("source")
    add("validate", cmd_validate, "validate all triples with SHACL")
    p = add("query", cmd_query, "run SPARQL (.rq file, name in kb/queries/, or inline text)")
    p.add_argument("sparql")
    p.add_argument("--param", action="append", help="value for {{name}} in the .rq (repeatable): --param frente=Infra")
    p.add_argument("--json", action="store_true")
    p.add_argument("--full-iri", action="store_true")
    p.add_argument("--width", type=int, default=80)
    p = add("show", cmd_show, "all claims about an entity, with source and precedence")
    p.add_argument("term", help="IRI, id:..., originalId, section number or exact label")
    p.add_argument("--limit", type=int, default=5)
    p.add_argument("--width", type=int, default=90)
    p = add("ask", cmd_ask, "natural-language question: most relevant entities with winning values and citation")
    p.add_argument("question")
    p.add_argument("--limit", type=int, default=6)
    p.add_argument("--width", type=int, default=300, help="value truncation width (excerpts use 3×)")
    p.add_argument("--source", "--fonte", dest="fonte", action="append",
                   help="restrict to sources whose path contains the text (repeatable)")
    p = add("search", cmd_search, "text search over labels/excerpts")
    p.add_argument("text")
    p.add_argument("--cls", help="restrict to a class, e.g. kb:Section")
    p.add_argument("--limit", type=int, default=40)
    p.add_argument("--width", type=int, default=70)
    add("wiki", cmd_wiki, "generate kb/wiki/ from the graph")
    p = add("stale", cmd_stale, "wiki/ notes with overdue review")
    p.add_argument("--all", action="store_true")
    p = add("new-note", cmd_new_note, "create wiki/YYYY-MM-DD-<slug>.md with review header (Zettelkasten)")
    p.add_argument("title")
    p.add_argument("--type", choices=["fleeting", "literature", "permanent"], default="fleeting")
    p.add_argument("--source", help="literature notes: relative path of the raw/ source")
    p = add("scaffold", cmd_scaffold, "create raw/ wiki/ kb/ docs/ and kb/config.yaml in the current folder")
    p.add_argument("--title")
    p.add_argument("--prefix", help="domain ontology prefix (e.g. prop, ctr)")
    p.add_argument("--language", default="en")
    p.add_argument("--versioning", choices=["git", "onedrive", "both"], default="git")
    p.add_argument("--force", action="store_true")
    p = add("migrate", cmd_migrate, "adopt a v1 workspace into the v2 engine (renames core terms; --dry-run first)")
    p.add_argument("--prefix")
    p.add_argument("--namespace")
    p.add_argument("--title")
    p.add_argument("--language", default="en")
    p.add_argument("--dry-run", action="store_true")
    p.add_argument("--force", action="store_true")
    p = add("log-cost", cmd_log_cost, "record tokens/time of an LLM build activity (from the subagent notification)")
    p.add_argument("--kind", required=True, choices=["images", "extract-llm", "graphify-build", "graphify-update", "answer-key", "other"])
    p.add_argument("--config", default="kb", choices=["kb", "graphify", "both"])
    p.add_argument("--tokens", type=int, required=True)
    p.add_argument("--duration-ms", type=int, default=0)
    p.add_argument("--tool-uses", type=int, default=0)
    p.add_argument("--note")
    p = add("costs", cmd_costs, "show logged build costs")
    p.add_argument("--json", action="store_true")
    p = add("vendor", cmd_vendor, "copy the engine into scripts/kb-engine/ for people without kb-setup (--check: compare only)")
    p.add_argument("--check", action="store_true", help="report absent | up to date | outdated; writes nothing")
    add("check-update", cmd_check_update, "check whether a newer kb-setup is published (network, 3 s timeout)")
    p = add("volume", cmd_volume, "corpus volume (files, pages, words, MB, estimated tokens) for evals")
    p.add_argument("path", nargs="?")
    p.add_argument("--json", action="store_true")
    args = ap.parse_args()
    if args.cmd in WRITE_CMDS and not getattr(args, "check", False) and (args.read_only or os.environ.get("KB_READONLY") == "1"):
        die(f"read-only mode: '{args.cmd}' writes to the workspace; ask the KB curator to run it")
    cfg = load_config()
    if args.cmd not in ("migrate", "scaffold"):
        for w in version_warnings(cfg):
            print(f"WARNING: {w}", file=sys.stderr)
    args.fn(args, cfg)


if __name__ == "__main__":
    main()
